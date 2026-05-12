# -*- coding: utf-8 -*-

import re
import sys
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent

DOCX_DIR = BASE_DIR / "docx_results"
TEMPLATE_XLSX = BASE_DIR / "template" / "스틸 챌린지 컬럼.xlsx"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_XLSX = OUTPUT_DIR / "스틸_챌린지_정리결과.xlsx"


SECTION_RUN = "Run Information"
SECTION_SIM = "Simulation Settings"
SECTION_COST = "Cost Breakdown"
SECTION_STEEL = "Steel Composition"
SECTION_RAW = "Raw Materials"
SECTION_ADD = "Additions"
SECTION_SLAG = "Slag Composition"

HEADER_WORDS = {
    "Element", "Current", "Min", "Max", "Name", "Qty [t]", "Qty[t]"
}

TEMPLATE_SECTIONS = {
    SECTION_RUN,
    SECTION_SIM,
    SECTION_COST,
    SECTION_STEEL,
    SECTION_RAW,
    SECTION_ADD,
    SECTION_SLAG,
    "Event Log",
}

KEY_ALIASES = {
    "status": "Status",
    "Time(in minutes)": "Time (in minutes)",
    "Time (in minutes)": "Time (in minutes)",
    "Intemal Low Alloyed": "Internal Low Alloyed",
    "Internal Low Alloyed": "Internal Low Alloyed",
    "Plate and structural": "Plate and Structural",
    "Plate and structural ": "Plate and Structural",
    "Plate and Structural": "Plate and Structural",
    "eafScrap Type05": "eafScrapType05",
    "eafScrapType05": "eafScrapType05",
    "eafScrap Type10": "eafScrapType10",
    "eafScrapType10": "eafScrapType10",
    "Turnnings": "Turnings",
    "Turnings": "Turnings",
}


def normalize_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_key(value):
    text = normalize_text(value)
    return KEY_ALIASES.get(text, text)


def clean_value(value):
    text = normalize_text(value)

    if text == "":
        return None

    text = text.replace("$", "")
    text = text.replace(",", "")
    text = text.replace("kWh/t", "")
    text = text.replace("kWh", "")
    text = text.replace("kg", "")
    text = re.sub(r"(?<=\d)t\b", "", text)
    text = text.strip()

    if re.fullmatch(r"[-+]?\d+(\.\d+)?", text):
        num = float(text)
        return int(num) if num.is_integer() else num

    return text


def table_rows(doc, table_index):
    if table_index >= len(doc.tables):
        return []

    rows = []

    for row in doc.tables[table_index].rows:
        cells = [normalize_text(cell.text) for cell in row.cells]

        if any(cells):
            rows.append(cells)

    return rows


def parse_key_value_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}
    total_energy_count = 0

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_key(row[0])

        if not key or key in HEADER_WORDS:
            continue

        value = clean_value(row[1])

        if section_name == SECTION_COST and key == "Total Energy":
            total_energy_count += 1
            save_key = "Total Energy" if total_energy_count == 1 else "Total Energy_2"
            data[(section_name, save_key)] = value
        else:
            data[(section_name, key)] = value

    return data


def parse_current_min_max_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_key(row[0])

        if not key or key in HEADER_WORDS:
            continue

        data[(section_name, key)] = clean_value(row[1])

    return data


def parse_event_log(doc):
    """
    Event Log 전체 문자열에서
    "00:00:00,이벤트내용" 형식을 하나씩 분리.
    """

    logs = []

    all_text_parts = []

    # 문단 텍스트 수집
    for p in doc.paragraphs:
        text = p.text
        if text:
            all_text_parts.append(text)

    # 표 안 텍스트도 수집
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                if text:
                    all_text_parts.append(text)

    full_text = "\n".join(all_text_parts)

    if "Event Log" not in full_text:
        return logs

    event_text = full_text.split("Event Log", 1)[1]

    # 따옴표 안의 로그만 하나씩 추출
    matches = re.findall(
        r'"(\d{2}:\d{2}:\d{2},.*?)"',
        event_text,
        flags=re.DOTALL
    )

    for item in matches:
        item = item.replace("\n", "")
        item = item.strip().rstrip(",")

        if "," not in item:
            continue

        time_part, event_part = item.split(",", 1)

        logs.append({
            "time": normalize_text(time_part),
            "event": unquote(normalize_text(event_part)),
        })

    return logs


def parse_docx(docx_path):
    doc = Document(docx_path)

    if len(doc.tables) < 7:
        raise ValueError(f"{docx_path.name} : table 개수 부족")

    data = {}

    data.update(parse_key_value_table(doc, 0, SECTION_RUN))
    data.update(parse_key_value_table(doc, 1, SECTION_SIM))
    data.update(parse_key_value_table(doc, 2, SECTION_COST))
    data.update(parse_current_min_max_table(doc, 3, SECTION_STEEL))
    data.update(parse_key_value_table(doc, 4, SECTION_RAW))
    data.update(parse_key_value_table(doc, 5, SECTION_ADD))
    data.update(parse_current_min_max_table(doc, 6, SECTION_SLAG))

    logs = parse_event_log(doc)

    return data, logs


def next_nonempty_value(ws, start_row):
    for r in range(start_row + 1, ws.max_row + 1):
        value = normalize_text(ws.cell(row=r, column=1).value)

        if value:
            return value

    return ""


def is_template_section_title(ws, row, value):
    if value not in TEMPLATE_SECTIONS:
        return False

    next_value = next_nonempty_value(ws, row)

    if value == SECTION_ADD:
        return next_value == "Name"

    if value == SECTION_RAW:
        return next_value == "Name"

    if value == SECTION_STEEL:
        return next_value == "Element"

    if value == SECTION_SLAG:
        return next_value == "Element"

    return True


def build_template_row_map(ws):
    row_map = {}
    current_section = None
    total_energy_count = 0

    for row in range(1, ws.max_row + 1):
        value = normalize_text(ws.cell(row=row, column=1).value)

        if not value:
            continue

        if is_template_section_title(ws, row, value):
            current_section = value
            total_energy_count = 0
            continue

        if value in HEADER_WORDS:
            continue

        if current_section is None:
            continue

        key = normalize_key(value)

        if current_section == SECTION_COST and key == "Total Energy":
            total_energy_count += 1
            key = "Total Energy" if total_energy_count == 1 else "Total Energy_2"

        row_map[(current_section, key)] = row

    return row_map


def clear_old_values(ws, start_col=2):
    for row in range(1, ws.max_row + 1):
        for col in range(start_col, ws.max_column + 1):
            ws.cell(row=row, column=col).value = None


def write_event_log_to_template_sheet(log_ws, all_results):
    """
    Sheet2에 Event1, Event2, Event3... 기준으로 로그를 하나씩 채움.

    A열: Event1, Event2, Event3...
    B열: Run 1 로그
    C열: Run 2 로그
    D열: Run 3 로그
    """

    # 기존 값 삭제
    for row in range(1, log_ws.max_row + 1):
        for col in range(2, log_ws.max_column + 1):
            log_ws.cell(row=row, column=col).value = None

    # 최대 로그 개수
    max_log_count = max(
        len(result_pack["logs"])
        for result_pack in all_results.values()
    )

    # 1행 Run 번호
    for col_idx, file_name in enumerate(all_results.keys(), start=2):
        log_ws.cell(row=1, column=col_idx).value = col_idx - 1

    # Event 행 생성 및 값 입력
    for log_idx in range(max_log_count):
        row_idx = log_idx + 2

        log_ws.cell(row=row_idx, column=1).value = f"Event{log_idx + 1}"

        for col_idx, (file_name, result_pack) in enumerate(all_results.items(), start=2):
            logs = result_pack["logs"]

            if log_idx < len(logs):
                time_text = logs[log_idx]["time"]
                event_text = logs[log_idx]["event"]

                log_ws.cell(
                    row=row_idx,
                    column=col_idx
                ).value = f"{time_text} | {event_text}"

    # 열 너비
    log_ws.column_dimensions["A"].width = 15

    for col in range(2, log_ws.max_column + 1):
        col_letter = log_ws.cell(row=1, column=col).column_letter
        log_ws.column_dimensions[col_letter].width = 90


def write_to_template(all_results):
    wb = load_workbook(TEMPLATE_XLSX)

    main_ws = wb.worksheets[0]

    if len(wb.worksheets) >= 2:
        log_ws = wb.worksheets[1]
    else:
        log_ws = wb.create_sheet("Event_Log")

    main_ws.title = "Result"
    log_ws.title = "Event_Log"

    clear_old_values(main_ws)

    row_map = build_template_row_map(main_ws)

    for col_idx, file_name in enumerate(all_results.keys(), start=2):
        main_ws.cell(row=1, column=col_idx).value = col_idx - 1

    unmatched_keys = {}

    for col_idx, (file_name, result_pack) in enumerate(all_results.items(), start=2):
        result = result_pack["data"]

        for section_key, value in result.items():
            target_row = row_map.get(section_key)

            if target_row is None:
                unmatched_keys.setdefault(file_name, []).append(section_key)
                continue

            main_ws.cell(row=target_row, column=col_idx).value = value

    write_event_log_to_template_sheet(log_ws, all_results)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)

    return unmatched_keys


def print_preview(all_results):
    check_keys = [
        (SECTION_RUN, "Status"),
        (SECTION_RUN, "Score"),
        (SECTION_SIM, "Steel Grade"),
        (SECTION_COST, "Time (in minutes)"),
        (SECTION_COST, "Cost Per Tonne"),
        (SECTION_STEEL, "C"),
        (SECTION_STEEL, "Cr"),
        (SECTION_RAW, "No1 Bundles"),
        (SECTION_RAW, "Direct Reduced Iron"),
        (SECTION_ADD, "Carbon"),
        (SECTION_ADD, "Chrome-Carbure Low S"),
        (SECTION_SLAG, "FeO"),
        (SECTION_SLAG, "Basicity"),
    ]

    print("\n[파싱 확인]")

    for file_name, result_pack in all_results.items():
        data = result_pack["data"]

        print(f"\n- {file_name}")

        for key in check_keys:
            print(f"{key[0]} > {key[1]} = {data.get(key)}")

        print(f"Event Log 개수 = {len(result_pack['logs'])}")


def main():
    if not TEMPLATE_XLSX.exists():
        print(f"[오류] 템플릿 없음:\n{TEMPLATE_XLSX}")
        sys.exit(1)

    if not DOCX_DIR.exists():
        print(f"[오류] DOCX 폴더 없음:\n{DOCX_DIR}")
        sys.exit(1)

    docx_files = sorted(
        f for f in DOCX_DIR.glob("*.docx")
        if not f.name.startswith("~$") and f.stat().st_size > 0
    )

    if not docx_files:
        print(f"[오류] DOCX 파일 없음:\n{DOCX_DIR}")
        sys.exit(1)

    all_results = {}

    for docx_path in docx_files:
        print(f"처리 중: {docx_path.name}")

        data, logs = parse_docx(docx_path)

        all_results[docx_path.name] = {
            "data": data,
            "logs": logs,
        }

    print_preview(all_results)

    unmatched = write_to_template(all_results)

    print("\n[완료]")
    print(OUTPUT_XLSX)

    if any(unmatched.values()):
        print("\n[템플릿에 없는 항목]")

        for file_name, keys in unmatched.items():
            if not keys:
                continue

            print(f"\n- {file_name}")

            for section, key in sorted(set(keys)):
                print(f"{section} > {key}")


if __name__ == "__main__":
    main()