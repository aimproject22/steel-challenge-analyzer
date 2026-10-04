"""Permission-neutral Excel generation helpers.

The workbook deliberately separates values reported by SteelUniversity from
values reconstructed or calculated by this application.  The wide internal
feature frame remains available for ML and diagnostics, but it is no longer
the public Excel schema.
"""

from __future__ import annotations

import math
import re
from io import BytesIO
from typing import Any, Iterable, Mapping, Optional

import numpy as np
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.linear_model import LinearRegression
from sklearn.model_selection import KFold, cross_val_score

from canonical import normalize_material_name
from export_features import export_token, extract_event_export_features


ILLEGAL_EXCEL_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]")
DUPLICATE_SUFFIX_RE = re.compile(r"_(\d+)$")
COMPOSITE_RUN_KEY_RE = re.compile(
    r"^Run Information > (?:User Id|Date|Status|Score)_\d+$"
)

CORE_META_COLUMNS = [
    "META_run_id", "META_run_date", "META_uploaded_at", "META_source",
    "META_source_run_index",
    "META_sender", "META_sender_email", "META_steel_user_id",
    "META_process_type", "META_uploader", "META_file_name", "META_status",
    "META_score", "META_user_level", "META_steel_grade",
]
CORE_PERF_COLUMNS = [
    "PERF_time_min", "PERF_time_sec", "PERF_tapping_mass_t",
    "PERF_tap_temperature_c", "PERF_total_energy_kwh",
    "PERF_energy_kwh_per_t_reported", "PERF_energy_kwh_per_t_calc",
    "PERF_energy_kwh_per_t_error", "PERF_score",
]
CORE_COST_COLUMNS = [
    "COST_source_power_raw", "COST_source_scrap_raw",
    "COST_source_additions_raw", "COST_source_other_consumables_raw",
    "COST_total_usd", "COST_per_tonne_usd",
    "COST_power_usd", "COST_scrap_usd", "COST_additions_usd",
    "COST_other_consumables_reconstructed_usd", "COST_other_total_usd",
    "COST_reconstructed_total_usd",
    "COST_reconciliation_error_usd", "COST_reconciliation_error_pct",
    "COST_source_scrap_is_time_duplicate", "COST_mapping_applied",
    "COST_mapping_type", "COST_mapping_confidence", "COST_mapping_warning",
    "COST_tapping_estimated_usd", "COST_oxygen_estimated_usd",
    "COST_carbon_injection_estimated_usd",
    "COST_electrode_breakage_estimated_usd", "COST_other_unexplained_usd",
    "COST_other_model_name", "COST_other_model_sample_count",
    "COST_other_model_cv_mae_usd", "COST_other_model_confidence",
    "COST_electrode_breakage_unit_estimate_usd",
    "COST_electrode_breakage_unit_mean_usd",
    "COST_electrode_breakage_unit_std_usd",
    "COST_electrode_breakage_unit_mad_usd",
    "COST_electrode_breakage_model_confidence",
]
CORE_QUALITY_COLUMNS = [
    "QUALITY_overall_pass", "QUALITY_full_validation_available",
    "QUALITY_chemistry_pass", "QUALITY_steel_pass",
    "QUALITY_steel_violation_count", "QUALITY_slag_pass",
    "QUALITY_slag_violation_count", "QUALITY_temperature_pass",
    "QUALITY_time_pass", "QUALITY_mass_pass", "QUALITY_co2_pass",
    "QUALITY_status_consistent", "QUALITY_status_matches_chemistry",
    "QUALITY_multi_run_merged_detected",
    "QUALITY_rm_source", "QUALITY_rm_parse_status", "QUALITY_rm_report_log_match",
    "QUALITY_rm_reconstruction_warning", "QUALITY_rm_total_report_t",
    "QUALITY_rm_total_log_t", "QUALITY_rm_total_difference_t",
    "QUALITY_add_source", "QUALITY_add_parse_status", "QUALITY_add_report_log_match",
    "QUALITY_add_reconstruction_warning", "QUALITY_add_total_report_kg",
    "QUALITY_add_total_log_kg", "QUALITY_add_total_difference_kg",
    "QUALITY_unknown_event_count", "QUALITY_warning_count",
    "QUALITY_warning_text",
]
GROUP_PREFIXES = (
    "RM_", "ADD_", "STEEL_", "SLAG_", "BASKET_", "PWR_", "O2_",
    "CINJ_", "ELEC_", "FAIL_", "EVTADD_", "ANALYSIS_", "TAP_",
    "DERIVED_", "RAW_",
)


def remove_illegal_characters(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    cleaned = ILLEGAL_EXCEL_RE.sub("", value)
    if cleaned.startswith(("=", "+", "-", "@")):
        cleaned = "'" + cleaned
    return cleaned


def clean_dataframe_for_excel(frame: pd.DataFrame) -> pd.DataFrame:
    cleaned = frame.copy()
    for column in cleaned.select_dtypes(include=["object", "string"]).columns:
        cleaned[column] = cleaned[column].map(remove_illegal_characters)
    return cleaned


def feature_frame(runs_df: pd.DataFrame) -> pd.DataFrame:
    """Retain the existing ML-facing helper for backward compatibility."""

    feature_columns = [
        column for column in runs_df.columns if str(column).startswith("feature >")
    ]
    identity_columns = [
        column
        for column in (
            "Run ID", "Run Date", "Sender", "Sender Email", "Steel User ID",
            "Steel Grade", "Score", "Cost Per Tonne", "Time",
        )
        if column in runs_df.columns
    ]
    return runs_df[identity_columns + feature_columns].copy()


def normalize_export_name(value: Any) -> str:
    return export_token(value)


def _is_missing(value: Any) -> bool:
    if value is None or value == "":
        return True
    if isinstance(value, (dict, list, tuple, set)):
        return False
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _value(row: Mapping[str, Any], *names: str) -> Any:
    for name in names:
        value = row.get(name)
        if not _is_missing(value):
            return value
    return None


def _number(value: Any) -> Optional[float]:
    if _is_missing(value):
        return None
    try:
        number = float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _run_key(value: Any) -> str:
    number = _number(value)
    if number is not None and number.is_integer():
        return str(int(number))
    return str(value or "")


def aggregate_logs_by_run(logs_df: pd.DataFrame) -> dict[str, list[dict[str, Any]]]:
    """Group logs once so export remains O(runs + logs)."""

    if logs_df is None or logs_df.empty or "run_id" not in logs_df.columns:
        return {}
    result: dict[str, list[dict[str, Any]]] = {}
    for run_id, group in logs_df.groupby("run_id", sort=False, dropna=False):
        result[_run_key(run_id)] = group.to_dict("records")
    return result


def _section_values(
    row: Mapping[str, Any], section: str
) -> tuple[dict[str, Any], list[str]]:
    """Collapse parser duplicate suffixes without trailing Excel columns."""

    prefix = f"{section} > "
    values: dict[str, Any] = {}
    warnings: list[str] = []
    for raw_key, value in row.items():
        key = str(raw_key)
        if not key.startswith(prefix) or " > " in key[len(prefix) :]:
            continue
        if isinstance(value, (dict, list)) or _is_missing(value):
            continue
        raw_name = key[len(prefix) :]
        name = DUPLICATE_SUFFIX_RE.sub("", raw_name)
        if name not in values:
            values[name] = value
        elif str(values[name]) != str(value):
            warnings.append(
                f"Conflicting duplicate {section} value for {name}: "
                f"kept {values[name]!r}, ignored {value!r}"
            )
    return values, warnings


def _normalize_section_materials(
    values: Mapping[str, Any], *, addition: bool
) -> tuple[dict[str, Any], list[str]]:
    normalized: dict[str, Any] = {}
    warnings: list[str] = []
    for raw_name, value in values.items():
        name = normalize_material_name(raw_name, addition=addition)
        if name not in normalized:
            normalized[name] = value
            continue
        previous = _number(normalized[name])
        current = _number(value)
        if previous is not None and current is not None:
            normalized[name] = previous + current
            warnings.append(f"Summed duplicate material {name}")
        elif str(normalized[name]) != str(value):
            warnings.append(
                f"Conflicting duplicate material {name}: kept {normalized[name]!r}"
            )
    return normalized, warnings


def build_material_columns(row: Mapping[str, Any]) -> tuple[dict[str, Any], list[str]]:
    values, warnings = _section_values(row, "Raw Materials")
    values, alias_warnings = _normalize_section_materials(values, addition=False)
    warnings.extend(alias_warnings)
    result: dict[str, Any] = {}
    numeric: list[float] = []
    nonzero = 0
    for name, value in values.items():
        number = _number(value)
        result[f"RM_{normalize_export_name(name)}_t"] = number if number is not None else value
        if number is not None:
            numeric.append(number)
            nonzero += int(number != 0)
    result["RM_total_mass_t"] = sum(numeric) if numeric else None
    result["RM_nonzero_material_count"] = nonzero if values else None
    return result, warnings


def build_addition_columns(row: Mapping[str, Any]) -> tuple[dict[str, Any], list[str]]:
    values, warnings = _section_values(row, "Additions")
    values, alias_warnings = _normalize_section_materials(values, addition=True)
    warnings.extend(alias_warnings)
    result: dict[str, Any] = {}
    numeric: list[float] = []
    nonzero = 0
    for name, value in values.items():
        number = _number(value)
        result[f"ADD_{normalize_export_name(name)}_kg"] = number if number is not None else value
        if number is not None:
            numeric.append(number)
            nonzero += int(number != 0)
    result["ADD_total_kg"] = sum(numeric) if numeric else None
    result["ADD_nonzero_count"] = nonzero if values else None
    return result, warnings


def _composition_source(
    row: Mapping[str, Any], section: str, nested_key: str
) -> dict[str, dict[str, Any]]:
    prefix = f"{section} > "
    result: dict[str, dict[str, Any]] = {}
    for raw_key, value in row.items():
        key = str(raw_key)
        if not key.startswith(prefix) or isinstance(value, (dict, list)):
            continue
        parts = key[len(prefix) :].split(" > ")
        element = parts[0]
        field = parts[-1].casefold() if len(parts) > 1 else "current"
        if field in {"current", "min", "max"} and not _is_missing(value):
            result.setdefault(element, {})[field] = value

    nested = row.get(nested_key)
    if isinstance(nested, Mapping):
        for element, values in nested.items():
            if not isinstance(values, Mapping):
                continue
            target = result.setdefault(str(element), {})
            for field in ("current", "min", "max"):
                if field not in target and not _is_missing(values.get(field)):
                    target[field] = values.get(field)
    return result


def _composition_columns(
    row: Mapping[str, Any], section: str, prefix: str, nested_key: str
) -> tuple[dict[str, Any], Optional[int], int]:
    source = _composition_source(row, section, nested_key)
    result: dict[str, Any] = {}
    pass_values: list[int] = []
    violations = 0
    for element, values in source.items():
        token = normalize_export_name(element)
        current = _number(values.get("current"))
        minimum = _number(values.get("min"))
        maximum = _number(values.get("max"))
        current_suffix = (
            "Basicity"
            if prefix == "SLAG_" and token.casefold() == "basicity"
            else (f"{token}_wt_pct" if prefix == "STEEL_" else f"{token}_pct")
        )
        result[f"{prefix}{current_suffix}"] = current
        has_spec = minimum is not None or maximum is not None
        if has_spec:
            result[f"{prefix}{token}_min"] = minimum
            result[f"{prefix}{token}_max"] = maximum
        if current is None or (minimum is None and maximum is None):
            within = None
        else:
            within = int(
                (minimum is None or current >= minimum)
                and (maximum is None or current <= maximum)
            )
            pass_values.append(within)
            violations += int(not within)
        if has_spec:
            result[f"{prefix}{token}_pass"] = within
            if minimum is not None:
                result[f"{prefix}{token}_margin_low"] = (
                    current - minimum if current is not None else None
                )
            if maximum is not None:
                result[f"{prefix}{token}_margin_high"] = (
                    maximum - current if current is not None else None
                )
            if minimum is not None and maximum is not None:
                result[f"{prefix}{token}_normalized_position"] = (
                    (current - minimum) / (maximum - minimum)
                    if current is not None and maximum != minimum else None
                )
    overall = int(all(pass_values)) if pass_values else None
    return result, overall, violations


def build_steel_columns(row: Mapping[str, Any]) -> tuple[dict[str, Any], Optional[int], int]:
    return _composition_columns(row, "Steel Composition", "STEEL_", "steel_composition")


def build_slag_columns(row: Mapping[str, Any]) -> tuple[dict[str, Any], Optional[int], int]:
    result, overall, violations = _composition_columns(
        row, "Slag Composition", "SLAG_", "slag_composition"
    )
    cao = _number(result.get("SLAG_CaO_pct"))
    sio2 = _number(result.get("SLAG_SiO2_pct"))
    reported = _number(result.get("SLAG_Basicity"))
    calculated = cao / sio2 if cao is not None and sio2 not in (None, 0) else None
    result["SLAG_Basicity_calc"] = calculated
    result["SLAG_Basicity_error"] = (
        reported - calculated if reported is not None and calculated is not None else None
    )
    return result, overall, violations


def build_event_export_features(
    logs: Iterable[Mapping[str, Any]], row: Mapping[str, Any]
) -> dict[str, Any]:
    return extract_event_export_features(
        logs,
        reported_time_minutes=_value(row, "Time", "Cost Breakdown > Time (in minutes)"),
        reported_energy_kwh=_value(row, "Total Energy kWh", "Cost Breakdown > Total Energy"),
    )


def _base_export_row(row: Mapping[str, Any]) -> dict[str, Any]:
    process_time = _number(_value(row, "Time", "Cost Breakdown > Time (in minutes)"))
    tapping_mass = _number(_value(row, "Tapping Mass", "Cost Breakdown > Tapping mass"))
    total_energy = _number(_value(row, "Total Energy kWh", "Cost Breakdown > Total Energy"))
    reported_energy_per_t = _number(
        _value(row, "Energy kWh/t", "Cost Breakdown > Total Energy_2")
    )
    calculated_energy_per_t = (
        total_energy / tapping_mass
        if total_energy is not None and tapping_mass not in (None, 0) else None
    )
    total_cost = _number(_value(row, "Total Cost", "Cost Breakdown > Total Cost"))
    cost_per_tonne = _number(_value(row, "Cost Per Tonne", "Cost Breakdown > Cost Per Tonne"))
    result = {
        "META_run_id": _value(row, "Run ID", "id"),
        "META_run_date": _value(row, "Run Date", "Run Information > Date"),
        "META_uploaded_at": _value(row, "Uploaded At", "uploaded_at", "created_at"),
        "META_source": _value(row, "Source", "source"),
        "META_source_run_index": _number(
            _value(row, "Source Run Index", "source_run_index")
        ),
        "META_sender": _value(row, "Sender", "sender_name"),
        "META_sender_email": _value(row, "Sender Email", "sender_email"),
        "META_steel_user_id": _value(row, "Steel User ID", "Run Information > User Id"),
        "META_process_type": _value(row, "Process Type", "Run Information > Process Type"),
        "META_uploader": _value(row, "Uploader", "uploader"),
        "META_file_name": _value(row, "File Name", "file_name"),
        "META_status": _number(_value(row, "Status", "Run Information > Status")),
        "META_score": _number(_value(row, "Score", "Run Information > Score")),
        "META_user_level": _value(row, "User Level", "Simulation Settings > User Level"),
        "META_steel_grade": _value(row, "Steel Grade", "Simulation Settings > Steel Grade"),
        "PERF_time_min": process_time,
        "PERF_time_sec": process_time * 60.0 if process_time is not None else None,
        "PERF_tapping_mass_t": tapping_mass,
        "PERF_tap_temperature_c": _number(_value(row, "Tap Temperature", "Cost Breakdown > Tap temperature")),
        "PERF_total_energy_kwh": total_energy,
        "PERF_energy_kwh_per_t_reported": reported_energy_per_t,
        "PERF_energy_kwh_per_t_calc": calculated_energy_per_t,
        "PERF_energy_kwh_per_t_error": (
            reported_energy_per_t - calculated_energy_per_t
            if reported_energy_per_t is not None and calculated_energy_per_t is not None else None
        ),
        "PERF_score": _number(_value(row, "Score", "Run Information > Score")),
        "COST_source_power_raw": _number(_value(row, "Cost Breakdown > Power")),
        "COST_source_scrap_raw": _number(_value(row, "Cost Breakdown > Scrap")),
        "COST_source_additions_raw": _number(_value(row, "Cost Breakdown > Additions")),
        "COST_source_other_consumables_raw": _number(
            _value(row, "Cost Breakdown > Other consumables")
        ),
        "COST_total_usd": total_cost,
        "COST_per_tonne_usd": cost_per_tonne,
        "COST_power_usd": None,
        "COST_scrap_usd": None,
        "COST_additions_usd": None,
        "COST_other_consumables_reconstructed_usd": None,
        "COST_other_total_usd": None,
        "COST_reconstructed_total_usd": None,
        "COST_reconciliation_error_usd": None,
        "COST_reconciliation_error_pct": None,
        "QUALITY_rm_parse_status": _value(
            row, "Parser Diagnostics > Raw Materials Parse Status"
        ),
        "QUALITY_add_parse_status": _value(
            row, "Parser Diagnostics > Additions Parse Status"
        ),
    }
    result["DERIVED_energy_per_process_min"] = (
        total_energy / process_time if total_energy is not None and process_time not in (None, 0) else None
    )
    result["DERIVED_cost_per_process_min"] = (
        total_cost / process_time if total_cost is not None and process_time not in (None, 0) else None
    )
    result["DERIVED_cost_per_tapping_t"] = (
        total_cost / tapping_mass if total_cost is not None and tapping_mass not in (None, 0) else None
    )
    return result


def _nearly_equal(left: Any, right: Any, *, absolute: float = 0.01) -> bool:
    first = _number(left)
    second = _number(right)
    if first is None or second is None:
        return False
    return abs(first - second) <= max(absolute, 1e-4 * max(abs(first), abs(second)))


def apply_cost_validation(export_row: dict[str, Any]) -> list[str]:
    """Preserve source labels and add an EAF-only semantic cost layer.

    SteelUniversity's EAF result table is known to emit shifted labels.  The
    raw values remain untouched; only the separate semantic fields below are
    corrected.  No material or electricity unit price is assumed.
    """

    source_power = _number(export_row.get("COST_source_power_raw"))
    source_scrap = _number(export_row.get("COST_source_scrap_raw"))
    source_additions = _number(export_row.get("COST_source_additions_raw"))
    source_other = _number(export_row.get("COST_source_other_consumables_raw"))
    total = _number(export_row.get("COST_total_usd"))
    process_time = _number(export_row.get("PERF_time_min"))
    process_type = str(export_row.get("META_process_type") or "").casefold()
    is_eaf = "electric arc furnace" in process_type
    warnings: list[str] = []
    if all(
        value is None
        for value in (source_power, source_scrap, source_additions, source_other, total)
    ):
        export_row.update(
            {
                "COST_source_scrap_is_time_duplicate": None,
                "COST_mapping_applied": None,
                "COST_mapping_type": "MISSING",
                "COST_mapping_confidence": "MISSING",
                "COST_mapping_warning": None,
            }
        )
        return warnings
    time_collision = (
        source_scrap is not None and _nearly_equal(source_scrap, process_time)
    )
    corrected: dict[str, Any] = {
        "COST_source_scrap_is_time_duplicate": int(time_collision),
    }
    if is_eaf:
        required = (source_power, source_additions, source_other, total)
        if all(value is not None for value in required):
            residual = total - source_power - source_additions - source_other
            if _nearly_equal(residual, 0.0, absolute=0.1):
                residual = 0.0
            corrected.update(
                {
                    "COST_power_usd": source_other,
                    "COST_scrap_usd": source_additions,
                    "COST_additions_usd": source_power,
                    "COST_other_consumables_reconstructed_usd": residual,
                    "COST_other_total_usd": residual,
                    "COST_mapping_applied": 1,
                    "COST_mapping_type": "STEELUNIVERSITY_EAF_KNOWN_SHIFT",
                    "COST_mapping_confidence": (
                        "HIGH" if time_collision and residual >= -0.01 else "MEDIUM"
                    ),
                    "COST_mapping_warning": (
                        "Known SteelUniversity EAF source-label shift corrected in the "
                        "semantic layer; raw labels are preserved."
                        + (
                            " Source Scrap duplicates reported process time."
                            if time_collision
                            else " Source Scrap did not duplicate reported process time."
                        )
                    ),
                }
            )
            warnings.append("Known EAF cost label shift mapped to semantic costs")
            if residual < -0.01:
                warnings.append("Reconstructed Other consumables is negative")
        else:
            corrected.update(
                {
                    "COST_mapping_applied": 0,
                    "COST_mapping_type": "STEELUNIVERSITY_EAF_INCOMPLETE_SOURCE",
                    "COST_mapping_confidence": "UNRESOLVED",
                    "COST_mapping_warning": "EAF cost mapping requires all source cost fields and Total Cost",
                }
            )
            warnings.append(str(corrected["COST_mapping_warning"]))
    else:
        corrected.update(
            {
                "COST_power_usd": source_power,
                "COST_scrap_usd": source_scrap,
                "COST_additions_usd": source_additions,
                "COST_other_consumables_reconstructed_usd": source_other,
                "COST_other_total_usd": source_other,
                "COST_mapping_applied": 0,
                "COST_mapping_type": "SOURCE_LABELS_AS_REPORTED",
                "COST_mapping_confidence": "REPORT",
                "COST_mapping_warning": None,
            }
        )

    export_row.update(corrected)
    corrected_values = [
        _number(export_row.get(column))
        for column in (
            "COST_power_usd",
            "COST_scrap_usd",
            "COST_additions_usd",
            "COST_other_consumables_reconstructed_usd",
        )
    ]
    if all(value is not None for value in corrected_values):
        reconstructed = sum(value for value in corrected_values if value is not None)
        export_row["COST_reconstructed_total_usd"] = reconstructed
        if total is not None:
            error = reconstructed - total
            export_row["COST_reconciliation_error_usd"] = error
            export_row["COST_reconciliation_error_pct"] = (
                error / total * 100.0 if total != 0 else None
            )
    return warnings


def _model_feature_value(frame: pd.DataFrame, column: str) -> pd.Series:
    """Return numeric model input while distinguishing measured zero from missing."""

    source = frame[column] if column in frame else pd.Series(np.nan, index=frame.index)
    values = pd.to_numeric(source, errors="coerce")
    if column == "O2_integrated_Nm3" and "O2_used" in frame:
        values = values.mask(values.isna() & frame["O2_used"].eq(0), 0.0)
    if column == "CINJ_integrated_kg" and "CINJ_used" in frame:
        values = values.mask(values.isna() & frame["CINJ_used"].eq(0), 0.0)
    return values


def apply_other_consumables_decomposition(frame: pd.DataFrame) -> pd.DataFrame:
    """Estimate Other-consumable components only when the dataset identifies them.

    Candidate non-negative models are fitted on EAF runs without electrode breaks.
    If the sample is insufficient, component estimates stay blank and the full
    residual remains unexplained.  This deliberately avoids hard-coded prices.
    """

    if frame.empty or "COST_other_total_usd" not in frame:
        return frame
    result = frame.copy()
    target = pd.to_numeric(result["COST_other_total_usd"], errors="coerce")
    break_source = (
        result["ELEC_break_count"]
        if "ELEC_break_count" in result
        else pd.Series(0.0, index=result.index)
    )
    breaks = pd.to_numeric(break_source, errors="coerce").fillna(0)
    is_eaf = result.get("META_process_type", pd.Series("", index=result.index)).fillna("").astype(str).str.contains(
        "electric arc furnace", case=False, regex=False
    )
    oxygen = _model_feature_value(result, "O2_integrated_Nm3")
    carbon = _model_feature_value(result, "CINJ_integrated_kg")
    tapping_mass = pd.to_numeric(
        result["PERF_tapping_mass_t"]
        if "PERF_tapping_mass_t" in result
        else pd.Series(np.nan, index=result.index),
        errors="coerce",
    )
    tapping_duration = pd.to_numeric(
        result["TAP_duration_sec"]
        if "TAP_duration_sec" in result
        else pd.Series(np.nan, index=result.index),
        errors="coerce",
    )

    candidates = {
        "BASE_O2_C": [("base", pd.Series(1.0, index=result.index)), ("oxygen", oxygen), ("carbon", carbon)],
        "TAP_MASS_O2_C": [("tapping", tapping_mass), ("oxygen", oxygen), ("carbon", carbon)],
        "TAP_DURATION_O2_C": [("tapping", tapping_duration), ("oxygen", oxygen), ("carbon", carbon)],
    }
    best: Optional[dict[str, Any]] = None
    for name, features in candidates.items():
        model_frame = pd.DataFrame(
            {feature_name: values for feature_name, values in features},
            index=result.index,
        )
        eligible = is_eaf & breaks.eq(0) & target.ge(0) & target.notna()
        eligible &= model_frame.notna().all(axis=1)
        x = model_frame.loc[eligible]
        y = target.loc[eligible]
        if len(x) < max(8, len(features) + 3):
            continue
        model = LinearRegression(positive=True, fit_intercept=False)
        folds = min(5, len(x))
        splitter = KFold(n_splits=folds, shuffle=True, random_state=42)
        mae = float(
            -cross_val_score(
                model,
                x.to_numpy(dtype=float),
                y.to_numpy(dtype=float),
                cv=splitter,
                scoring="neg_mean_absolute_error",
            ).mean()
        )
        model.fit(x.to_numpy(dtype=float), y.to_numpy(dtype=float))
        candidate = {
            "name": name,
            "features": features,
            "columns": list(x.columns),
            "model": model,
            "mae": mae,
            "sample_count": len(x),
            "score": mae * (1.0 + 0.02 * len(features)),
        }
        if best is None or candidate["score"] < best["score"]:
            best = candidate

    component_columns = (
        "COST_tapping_estimated_usd",
        "COST_oxygen_estimated_usd",
        "COST_carbon_injection_estimated_usd",
        "COST_electrode_breakage_estimated_usd",
        "COST_other_unexplained_usd",
    )
    for column in component_columns:
        result[column] = np.nan
    if best is None:
        result["COST_other_unexplained_usd"] = target
        result.loc[is_eaf, "COST_other_model_name"] = "UNAVAILABLE"
        result.loc[is_eaf, "COST_other_model_sample_count"] = int(
            (is_eaf & breaks.eq(0) & target.notna()).sum()
        )
        result["COST_other_model_cv_mae_usd"] = np.nan
        result.loc[is_eaf, "COST_other_model_confidence"] = "INSUFFICIENT_DATA"
        result.loc[
            is_eaf, "COST_electrode_breakage_model_confidence"
        ] = "INSUFFICIENT_DATA"
        return result

    coefficients = dict(zip(best["columns"], best["model"].coef_))
    tapping_component = pd.Series(0.0, index=result.index)
    if best["name"] == "BASE_O2_C":
        tapping_component[:] = float(coefficients.get("base", 0.0))
    elif best["name"] == "TAP_MASS_O2_C":
        tapping_component = tapping_mass * float(coefficients.get("tapping", 0.0))
    else:
        tapping_component = tapping_duration * float(coefficients.get("tapping", 0.0))
    oxygen_component = oxygen * float(coefficients.get("oxygen", 0.0))
    carbon_component = carbon * float(coefficients.get("carbon", 0.0))
    base_total = tapping_component + oxygen_component + carbon_component

    break_eligible = is_eaf & breaks.gt(0) & target.notna() & base_total.notna()
    unit_estimates = ((target - base_total) / breaks).loc[break_eligible]
    unit_estimates = unit_estimates[unit_estimates.ge(0)]
    unit_median = float(unit_estimates.median()) if len(unit_estimates) else np.nan
    unit_mean = float(unit_estimates.mean()) if len(unit_estimates) else np.nan
    unit_std = float(unit_estimates.std(ddof=0)) if len(unit_estimates) else np.nan
    unit_mad = (
        float((unit_estimates - unit_median).abs().median())
        if len(unit_estimates)
        else np.nan
    )
    electrode_component = breaks * unit_median if math.isfinite(unit_median) else pd.Series(np.nan, index=result.index)
    electrode_component = electrode_component.where(breaks.gt(0), 0.0)
    explained = base_total + electrode_component
    valid_prediction = is_eaf & target.notna() & base_total.notna()
    confidence = (
        "HIGH" if best["sample_count"] >= 30 and best["mae"] <= max(5.0, float(target.median()) * 0.05)
        else "MEDIUM" if best["sample_count"] >= 15
        else "LOW"
    )
    electrode_confidence = (
        "HIGH" if len(unit_estimates) >= 5 and unit_mad <= max(5.0, abs(unit_median) * 0.15)
        else "MEDIUM" if len(unit_estimates) >= 3
        else "LOW" if len(unit_estimates) >= 1
        else "INSUFFICIENT_DATA"
    )
    result.loc[valid_prediction, "COST_tapping_estimated_usd"] = tapping_component
    result.loc[valid_prediction, "COST_oxygen_estimated_usd"] = oxygen_component
    result.loc[valid_prediction, "COST_carbon_injection_estimated_usd"] = carbon_component
    result.loc[valid_prediction, "COST_electrode_breakage_estimated_usd"] = electrode_component
    result.loc[valid_prediction, "COST_other_unexplained_usd"] = target - explained
    result.loc[is_eaf, "COST_other_model_name"] = best["name"]
    result.loc[is_eaf, "COST_other_model_sample_count"] = best["sample_count"]
    result.loc[is_eaf, "COST_other_model_cv_mae_usd"] = best["mae"]
    result.loc[is_eaf, "COST_other_model_confidence"] = confidence
    result.loc[is_eaf, "COST_electrode_breakage_unit_estimate_usd"] = unit_median
    result.loc[is_eaf, "COST_electrode_breakage_unit_mean_usd"] = unit_mean
    result.loc[is_eaf, "COST_electrode_breakage_unit_std_usd"] = unit_std
    result.loc[is_eaf, "COST_electrode_breakage_unit_mad_usd"] = unit_mad
    result.loc[
        is_eaf, "COST_electrode_breakage_model_confidence"
    ] = electrode_confidence
    return result


def _detail_values(
    values: Mapping[str, Any], prefix: str, suffix: str, excluded: set[str]
) -> dict[str, float]:
    result: dict[str, float] = {}
    for column, value in values.items():
        if not str(column).startswith(prefix) or not str(column).endswith(suffix):
            continue
        if column in excluded:
            continue
        number = _number(value)
        if number is not None:
            token = str(column)[len(prefix) : -len(suffix)]
            result[token] = result.get(token, 0.0) + number
    return result


def _material_match(report: Mapping[str, float], logged: Mapping[str, float]) -> bool:
    if not report or not logged:
        return False
    for name in set(report) | set(logged):
        if not _nearly_equal(report.get(name, 0.0), logged.get(name, 0.0), absolute=0.01):
            return False
    return True


def apply_material_fallbacks(export_row: dict[str, Any]) -> list[str]:
    """Prefer report tables, otherwise reconstruct material totals from events."""

    warnings: list[str] = []
    report_rm = _detail_values(
        export_row,
        "RM_",
        "_t",
        {"RM_total_mass_t"},
    )
    logged_rm = _detail_values(export_row, "BASKET_", "_total_t", set())
    report_add = _detail_values(
        export_row,
        "ADD_",
        "_kg",
        {"ADD_total_kg"},
    )
    logged_add = _detail_values(export_row, "EVTADD_", "_total_kg", set())

    rm_report_total = sum(report_rm.values()) if report_rm else None
    rm_log_total = sum(logged_rm.values()) if logged_rm else None
    add_report_total = sum(report_add.values()) if report_add else None
    add_log_total = sum(logged_add.values()) if logged_add else None
    export_row.update(
        {
            "QUALITY_rm_total_report_t": rm_report_total,
            "QUALITY_rm_total_log_t": rm_log_total,
            "QUALITY_rm_total_difference_t": (
                rm_report_total - rm_log_total
                if rm_report_total is not None and rm_log_total is not None else None
            ),
            "QUALITY_add_total_report_kg": add_report_total,
            "QUALITY_add_total_log_kg": add_log_total,
            "QUALITY_add_total_difference_kg": (
                add_report_total - add_log_total
                if add_report_total is not None and add_log_total is not None else None
            ),
        }
    )

    if report_rm:
        export_row["QUALITY_rm_parse_status"] = "REPORT_PARSED"
        match = _material_match(report_rm, logged_rm) if logged_rm else None
        export_row["QUALITY_rm_source"] = (
            "REPORT_AND_LOG_VERIFIED" if match else "REPORT"
        )
        export_row["QUALITY_rm_report_log_match"] = int(match) if match is not None else None
        if match is False:
            export_row["QUALITY_rm_reconstruction_warning"] = "Report and Event Log RM differ"
            warnings.append("Raw Materials report/log mismatch")
    elif logged_rm:
        for token, value in logged_rm.items():
            export_row[f"RM_{token}_t"] = value
        export_row["RM_total_mass_t"] = rm_log_total
        export_row["RM_nonzero_material_count"] = sum(value != 0 for value in logged_rm.values())
        export_row["QUALITY_rm_source"] = "EVENT_LOG_RECONSTRUCTED"
        export_row["QUALITY_rm_parse_status"] = "EVENT_RECONSTRUCTED"
        export_row["QUALITY_rm_reconstruction_warning"] = "Report RM missing; reconstructed from Event Log"
    else:
        export_row["QUALITY_rm_source"] = "MISSING"
        if not export_row.get("QUALITY_rm_parse_status"):
            export_row["QUALITY_rm_parse_status"] = "MISSING"
        export_row["QUALITY_rm_reconstruction_warning"] = "Raw Materials unavailable"
        warnings.append("Raw Materials missing")

    if report_add:
        export_row["QUALITY_add_parse_status"] = "REPORT_PARSED"
        match = _material_match(report_add, logged_add) if logged_add else None
        export_row["QUALITY_add_source"] = (
            "REPORT_AND_LOG_VERIFIED" if match else "REPORT"
        )
        export_row["QUALITY_add_report_log_match"] = int(match) if match is not None else None
        if match is False:
            export_row["QUALITY_add_reconstruction_warning"] = "Report and Event Log additions differ"
            warnings.append("Additions report/log mismatch")
    elif logged_add:
        for token, value in logged_add.items():
            export_row[f"ADD_{token}_kg"] = value
        export_row["ADD_total_kg"] = add_log_total
        export_row["ADD_nonzero_count"] = sum(value != 0 for value in logged_add.values())
        export_row["QUALITY_add_source"] = "EVENT_LOG_RECONSTRUCTED"
        export_row["QUALITY_add_parse_status"] = "EVENT_RECONSTRUCTED"
        export_row["QUALITY_add_reconstruction_warning"] = "Report additions missing; reconstructed from Event Log"
    else:
        export_row["QUALITY_add_source"] = "MISSING"
        if not export_row.get("QUALITY_add_parse_status"):
            export_row["QUALITY_add_parse_status"] = "MISSING"
        export_row["QUALITY_add_reconstruction_warning"] = "Additions unavailable"
    return warnings


def detect_multi_run_merged(row: Mapping[str, Any], logs: Iterable[Mapping[str, Any]]) -> bool:
    duplicate_fields = any(
        COMPOSITE_RUN_KEY_RE.match(str(key)) and not _is_missing(value)
        for key, value in row.items()
    )
    ordered = sorted(
        list(logs),
        key=lambda item: int(_number(item.get("log_no")) or 0),
    )
    events = [str(item.get("event") or "").casefold() for item in ordered]
    tapping_seen = False
    clock_reset = False
    tapping_seconds: Optional[float] = None
    for item, event in zip(ordered, events):
        seconds = _number(item.get("event_seconds"))
        if seconds is None:
            time_text = str(item.get("event_time") or item.get("time") or "")
            parts = time_text.split(":")
            try:
                seconds = float(parts[0]) * 3600 + float(parts[1]) * 60 + float(parts[2])
            except (IndexError, TypeError, ValueError):
                seconds = None
        if tapping_seen and seconds is not None and tapping_seconds is not None:
            if seconds < tapping_seconds:
                clock_reset = True
        if "tapping complete" in event:
            tapping_seen = True
            tapping_seconds = seconds
    return bool(
        duplicate_fields
        or sum("tapping complete" in event for event in events) > 1
        or sum("selected steel grade" in event for event in events) > 1
        or clock_reset
    )


def drop_all_null_export_columns(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    if frame.empty:
        return frame, []
    null_columns = [
        column
        for column in frame.columns
        if frame[column].map(_is_missing).all()
    ]
    return frame.drop(columns=null_columns), null_columns


def drop_redundant_export_columns(frame: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Drop only known presentation duplicates that are identical for all runs."""

    pairs = (
        ("PERF_score", "META_score"),
        ("PWR_energy_reported_kwh", "PERF_total_energy_kwh"),
        ("META_uploader", "META_sender"),
    )
    dropped: list[str] = []
    for candidate, retained in pairs:
        if candidate in frame and retained in frame and frame[candidate].equals(frame[retained]):
            dropped.append(candidate)
    return frame.drop(columns=dropped), dropped


def _column_order(columns: Iterable[str]) -> list[str]:
    available = set(columns)
    ordered: list[str] = []

    def natural_key(value: str) -> list[tuple[int, Any]]:
        return [
            (0, int(part)) if part.isdigit() else (1, part.casefold())
            for part in re.split(r"(\d+)", value)
            if part
        ]

    for column in CORE_META_COLUMNS + CORE_PERF_COLUMNS + CORE_COST_COLUMNS + CORE_QUALITY_COLUMNS:
        if column in available:
            ordered.append(column)
    for prefix in GROUP_PREFIXES:
        group = [column for column in available if column.startswith(prefix)]
        if prefix == "RAW_":
            preferred = ["RAW_event_sequence", "RAW_unknown_events"]
            ordered.extend(column for column in preferred if column in group)
            group = [column for column in group if column not in preferred]
        ordered.extend(sorted(group, key=natural_key))
    ordered.extend(sorted(available - set(ordered), key=lambda value: value.casefold()))
    return ordered


def build_all_runs_export_df(
    runs_df: pd.DataFrame,
    logs_df: pd.DataFrame,
    ml_features_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """Return the deterministic 1-run-per-row ALL_RUNS export frame."""

    del ml_features_df
    log_groups = aggregate_logs_by_run(logs_df)
    records: list[dict[str, Any]] = []
    source_runs = runs_df if runs_df is not None else pd.DataFrame()
    source_records = source_runs.to_dict("records")
    for row in source_records:
        export_row = _base_export_row(row)
        warnings: list[str] = []
        warnings.extend(apply_cost_validation(export_row))
        materials, material_warnings = build_material_columns(row)
        additions, addition_warnings = build_addition_columns(row)
        steel, steel_pass, steel_violations = build_steel_columns(row)
        slag, slag_pass, slag_violations = build_slag_columns(row)
        warnings.extend(material_warnings)
        warnings.extend(addition_warnings)
        export_row.update(materials)
        export_row.update(additions)
        export_row.update(steel)
        export_row.update(slag)

        run_logs = log_groups.get(_run_key(export_row.get("META_run_id")), [])
        event_features = build_event_export_features(run_logs, row)
        warnings.extend(event_features.pop("_warnings", []))
        export_row.update(event_features)
        warnings.extend(apply_material_fallbacks(export_row))
        merged = int(detect_multi_run_merged(row, run_logs))
        if merged:
            warnings.append("Multiple simulation runs may be merged in this DB record")
        export_row.update({
            "QUALITY_steel_pass": steel_pass,
            "QUALITY_steel_violation_count": steel_violations,
            "QUALITY_slag_pass": slag_pass,
            "QUALITY_slag_violation_count": slag_violations,
            "QUALITY_temperature_pass": None,
            "QUALITY_time_pass": None,
            "QUALITY_mass_pass": None,
            "QUALITY_co2_pass": None,
            "QUALITY_full_validation_available": 0,
            "QUALITY_multi_run_merged_detected": merged,
        })
        known_quality = [value for value in (steel_pass, slag_pass) if value is not None]
        chemistry = int(all(known_quality)) if known_quality else None
        export_row["QUALITY_chemistry_pass"] = chemistry
        export_row["QUALITY_overall_pass"] = None
        status = _number(export_row.get("META_status"))
        export_row["QUALITY_status_consistent"] = None
        export_row["QUALITY_status_matches_chemistry"] = (
            int(status == chemistry)
            if status is not None and chemistry is not None else None
        )
        if not run_logs:
            warnings.append("Event Log missing")
        export_row["QUALITY_warning_count"] = len(warnings)
        export_row["QUALITY_warning_text"] = " | ".join(warnings)
        records.append(export_row)

    frame = pd.DataFrame(records)
    if frame.empty:
        frame = pd.DataFrame(
            columns=CORE_META_COLUMNS + CORE_PERF_COLUMNS + CORE_COST_COLUMNS + CORE_QUALITY_COLUMNS
        )
        return frame
    frame = apply_other_consumables_decomposition(frame)
    frame, _ = drop_all_null_export_columns(frame)
    frame, _ = drop_redundant_export_columns(frame)
    return frame.reindex(columns=_column_order(frame.columns))


RESULTS_CORE_COLUMNS = [
    "Run_ID", "User_ID", "Run_Date", "Status", "Score", "User_Level",
    "Steel_Grade", "Process_Type", "Source", "Source_Run_Index", "Sender",
    "Sender_Email", "Time_min", "Tapping_Mass_t", "Tap_Temperature_C",
    "Total_Energy_kWh", "Energy_kWh_per_t", "Reported_Power",
    "Reported_Scrap", "Reported_Additions", "Reported_Other_Consumables",
    "Total_Cost", "Cost_Per_Tonne",
]


def _reported_composition_columns(
    row: Mapping[str, Any], section: str, label: str, nested_key: str
) -> dict[str, Any]:
    """Return only Current/Min/Max values present in the source result."""

    result: dict[str, Any] = {}
    for element, values in _composition_source(row, section, nested_key).items():
        token = normalize_export_name(element)
        for source_name, suffix in (("current", ""), ("min", "_Min"), ("max", "_Max")):
            value = values.get(source_name)
            if not _is_missing(value):
                number = _number(value)
                result[f"{label}_{token}{suffix}"] = number if number is not None else value
    return result


def _natural_column_key(value: str) -> list[tuple[int, Any]]:
    return [
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", value)
        if part
    ]


def build_results_dataframe(
    runs_df: pd.DataFrame, logs_df: pd.DataFrame
) -> pd.DataFrame:
    """Build the human-readable source/report sheet without reconstruction."""

    log_groups = aggregate_logs_by_run(logs_df)
    records: list[dict[str, Any]] = []
    for row in (runs_df if runs_df is not None else pd.DataFrame()).to_dict("records"):
        base = _base_export_row(row)
        record: dict[str, Any] = {
            "Run_ID": base.get("META_run_id"),
            "User_ID": base.get("META_steel_user_id"),
            "Run_Date": base.get("META_run_date"),
            "Status": base.get("META_status"),
            "Score": base.get("META_score"),
            "User_Level": base.get("META_user_level"),
            "Steel_Grade": base.get("META_steel_grade"),
            "Process_Type": base.get("META_process_type"),
            "Source": base.get("META_source"),
            "Source_Run_Index": base.get("META_source_run_index"),
            "Sender": base.get("META_sender"),
            "Sender_Email": base.get("META_sender_email"),
            "Time_min": base.get("PERF_time_min"),
            "Tapping_Mass_t": base.get("PERF_tapping_mass_t"),
            "Tap_Temperature_C": base.get("PERF_tap_temperature_c"),
            "Total_Energy_kWh": base.get("PERF_total_energy_kwh"),
            "Energy_kWh_per_t": base.get("PERF_energy_kwh_per_t_reported"),
            "Reported_Power": base.get("COST_source_power_raw"),
            "Reported_Scrap": base.get("COST_source_scrap_raw"),
            "Reported_Additions": base.get("COST_source_additions_raw"),
            "Reported_Other_Consumables": base.get("COST_source_other_consumables_raw"),
            "Total_Cost": base.get("COST_total_usd"),
            "Cost_Per_Tonne": base.get("COST_per_tonne_usd"),
        }
        # These helpers inspect only the report tables.  Their summary fields
        # are intentionally excluded so RESULTS contains only reported items.
        materials, _ = build_material_columns(row)
        additions, _ = build_addition_columns(row)
        record.update(
            (key, value)
            for key, value in materials.items()
            if key not in {"RM_total_mass_t", "RM_nonzero_material_count"}
        )
        record.update(
            (key, value)
            for key, value in additions.items()
            if key not in {"ADD_total_kg", "ADD_nonzero_count"}
        )
        record.update(
            _reported_composition_columns(
                row, "Steel Composition", "Steel", "steel_composition"
            )
        )
        record.update(
            _reported_composition_columns(
                row, "Slag Composition", "Slag", "slag_composition"
            )
        )
        run_logs = log_groups.get(_run_key(record.get("Run_ID")), [])
        raw = build_event_export_features(run_logs, row)
        raw_columns = sorted(
            (key for key in raw if key.startswith("RAW_event_log_")),
            key=_natural_column_key,
        )
        for key in raw_columns:
            record[key.replace("RAW_event_log_", "Event_Log_Raw_")] = raw[key]
        records.append(record)

    frame = pd.DataFrame(records)
    if frame.empty:
        return pd.DataFrame(columns=["Run_ID"])
    frame, _ = drop_all_null_export_columns(frame)
    groups = ("RM_", "ADD_", "Steel_", "Slag_", "Event_Log_Raw_")
    ordered = [column for column in RESULTS_CORE_COLUMNS if column in frame]
    for prefix in groups:
        ordered.extend(
            sorted(
                (
                    column
                    for column in frame
                    if column.startswith(prefix) and column not in ordered
                ),
                key=_natural_column_key,
            )
        )
    ordered.extend(column for column in frame if column not in ordered)
    return frame.reindex(columns=ordered)


ANALYSIS_FIELD_MAP = {
    "COST_additions_usd": "Corrected_Additions_Cost_USD",
    "COST_scrap_usd": "Corrected_Scrap_Cost_USD",
    "COST_power_usd": "Corrected_Power_Cost_USD",
    "COST_mapping_applied": "Cost_Mapping_Applied",
    "COST_mapping_type": "Cost_Mapping_Type",
    "COST_source_scrap_is_time_duplicate": "Reported_Scrap_Equals_Time",
    "QUALITY_rm_source": "RM_Source",
    "QUALITY_rm_report_log_match": "RM_Report_Log_Match",
    "QUALITY_rm_total_log_t": "RM_Total_Event_t",
    "QUALITY_add_source": "ADD_Source",
    "QUALITY_add_report_log_match": "ADD_Report_Log_Match",
    "QUALITY_add_total_log_kg": "ADD_Total_Event_kg",
    "BASKET_count": "Basket_Count",
    "BASKET_total_charged_t": "Basket_Total_Charged_t",
    "PWR_first_on_sec": "Power_First_On_sec",
    "PWR_last_change_sec": "Power_Last_Change_sec",
    "PWR_max_mw": "Power_Max_MW",
    "PWR_min_nonzero_mw": "Power_Min_Nonzero_MW",
    "PWR_change_count": "Power_Change_Count",
    "PWR_total_on_sec": "Power_Total_On_sec",
    "PWR_total_off_sec": "Power_Total_Off_sec",
    "PWR_time_weighted_avg_mw": "Power_Time_Weighted_Avg_MW",
    "PWR_setpoint_integral_MW_min": "Power_Setpoint_Integral_MW_min",
    "PWR_setpoint_equivalent_MWh": "Power_Setpoint_Equivalent_MWh",
    "O2_used": "Oxygen_Used",
    "O2_first_on_sec": "Oxygen_First_On_sec",
    "O2_last_off_sec": "Oxygen_Last_Off_sec",
    "O2_total_on_sec": "Oxygen_Total_On_sec",
    "O2_max_flow_Nm3_min": "Oxygen_Max_Flow_Nm3_min",
    "O2_integrated_Nm3": "Oxygen_Integrated_Nm3",
    "O2_change_count": "Oxygen_Change_Count",
    "CINJ_used": "Carbon_Injection_Used",
    "CINJ_total_on_sec": "Carbon_Injection_Total_On_sec",
    "CINJ_max_flow": "Carbon_Injection_Max_kg_min",
    "CINJ_integrated_kg": "Carbon_Injection_Integrated_kg",
    "CINJ_change_count": "Carbon_Injection_Change_Count",
    "ELEC_break_count": "Electrode_Break_Count",
    "ELEC_replacement_count": "Electrode_Replacement_Count",
    "ELEC_first_break_sec": "Electrode_First_Break_sec",
    "ELEC_first_replacement_sec": "Electrode_First_Replacement_sec",
    "ELEC_total_downtime_sec": "Electrode_Total_Downtime_sec",
    "ANALYSIS_request_count": "Sample_Request_Count",
    "ANALYSIS_received_count": "Sample_Result_Count",
    "ANALYSIS_first_request_sec": "First_Sample_Request_sec",
    "ANALYSIS_first_received_sec": "First_Sample_Result_sec",
    "ANALYSIS_avg_wait_sec": "Average_Analysis_Wait_sec",
    "ANALYSIS_max_wait_sec": "Max_Analysis_Wait_sec",
    "ANALYSIS_total_wait_sec": "Total_Analysis_Wait_sec",
    "TAP_start_sec": "Tapping_Start_sec",
    "TAP_complete_sec": "Tapping_Complete_sec",
    "TAP_duration_sec": "Tapping_Duration_sec",
    "EVTADD_first_addition_sec": "First_Addition_sec",
    "EVTADD_last_addition_sec": "Last_Addition_sec",
    "EVTADD_event_count": "Addition_Event_Count",
    "EVTADD_High_C_Ferro_Manganese_first_sec": "HC_FeMn_First_Addition_sec",
    "EVTADD_Lime_first_sec": "Lime_First_Addition_sec",
    "EVTADD_Dolomite_first_sec": "Dolomite_First_Addition_sec",
    "EVTADD_Iron_Oxide_first_sec": "Iron_Oxide_First_Addition_sec",
    "QUALITY_steel_pass": "Steel_Pass",
    "QUALITY_slag_pass": "Slag_Pass",
    "QUALITY_chemistry_pass": "Chemistry_Pass",
    "QUALITY_time_pass": "Time_Pass",
    "QUALITY_temperature_pass": "Temperature_Pass",
    "QUALITY_mass_pass": "Mass_Pass",
    "QUALITY_co2_pass": "CO2_Pass",
    "QUALITY_overall_pass": "Overall_Pass",
    "QUALITY_full_validation_available": "Validation_Available",
    "QUALITY_unknown_event_count": "Unknown_Event_Count",
    "QUALITY_warning_count": "Warning_Count",
    "QUALITY_warning_text": "Warning_Text",
    "QUALITY_multi_run_merged_detected": "Multi_Run_Merged_Suspect",
}


def _analysis_cost_fields(row: Mapping[str, Any]) -> dict[str, Any]:
    mapping_applied = int(_number(row.get("COST_mapping_applied")) or 0) == 1
    raw_residual: Optional[float] = None
    if mapping_applied:
        total = _number(row.get("COST_total_usd"))
        additions = _number(row.get("COST_source_power_raw"))
        scrap = _number(row.get("COST_source_additions_raw"))
        power = _number(row.get("COST_source_other_consumables_raw"))
        if all(value is not None for value in (total, additions, scrap, power)):
            raw_residual = float(total - additions - scrap - power)
    else:
        raw_residual = _number(row.get("COST_other_consumables_reconstructed_usd"))
    rounding_adjustment = (
        raw_residual if raw_residual is not None and abs(raw_residual) <= 0.5 else 0.0
    )
    other = 0.0 if raw_residual is not None and abs(raw_residual) <= 0.5 else raw_residual
    corrected = [
        _number(row.get("COST_additions_usd")),
        _number(row.get("COST_scrap_usd")),
        _number(row.get("COST_power_usd")),
        other,
    ]
    total = _number(row.get("COST_total_usd"))
    reconciliation = (
        sum(value for value in corrected if value is not None) - total
        if total is not None and all(value is not None for value in corrected)
        else None
    )
    return {
        "Other_Consumables_Raw_Residual": raw_residual,
        "Other_Consumables_USD": other,
        "Cost_Rounding_Adjustment_USD": rounding_adjustment,
        "Cost_Reconciliation_Error_USD": reconciliation,
    }


def _analysis_other_cost_fields(
    record: dict[str, Any], source: Mapping[str, Any]
) -> None:
    """Export a conservative, auditable Other Consumables decomposition.

    Electrode events establish that an incident occurred, but they do not
    establish a monetary replacement cost.  Until a verified electrode price
    source is available, any remaining cost stays explicitly unexplained.
    """

    other_total = _number(record.get("Other_Consumables_USD"))
    model_confidence = str(source.get("COST_other_model_confidence") or "")
    component_fields = (
        ("COST_oxygen_estimated_usd", "Oxygen_Cost_Estimated_USD"),
        ("COST_carbon_injection_estimated_usd", "Carbon_Cost_Estimated_USD"),
        ("COST_tapping_estimated_usd", "Tapping_Cost_Estimated_USD"),
    )

    known_components: list[float] = []
    if model_confidence in {"MEDIUM", "HIGH"}:
        for source_name, export_name in component_fields:
            value = _number(source.get(source_name))
            if value is not None:
                record[export_name] = value
                known_components.append(value)
        record["Other_Cost_Model_Confidence"] = model_confidence

    if other_total is None:
        return

    raw_unexplained = other_total - sum(known_components)
    component_rounding_adjustment = (
        raw_unexplained if abs(raw_unexplained) <= 0.5 else 0.0
    )
    unexplained = 0.0 if component_rounding_adjustment else raw_unexplained

    # Reuse the existing rounding column when it is not already carrying a
    # top-level cost-mapping adjustment.
    existing_rounding = _number(record.get("Cost_Rounding_Adjustment_USD"))
    if component_rounding_adjustment and (
        existing_rounding is None or abs(existing_rounding) < 1e-12
    ):
        record["Cost_Rounding_Adjustment_USD"] = component_rounding_adjustment

    record["Other_Unexplained_USD"] = unexplained
    record["Other_Cost_Reconciliation_Error_USD"] = (
        other_total
        - sum(known_components)
        - unexplained
    )


def build_analysis_dataframe(
    runs_df: pd.DataFrame,
    logs_df: pd.DataFrame,
    ml_features_df: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """Build a compact calculated/reconstructed sheet from the internal frame."""

    wide = build_all_runs_export_df(runs_df, logs_df, ml_features_df)
    records: list[dict[str, Any]] = []
    for source in wide.to_dict("records"):
        record: dict[str, Any] = {"Run_ID": source.get("META_run_id")}
        for old, new in ANALYSIS_FIELD_MAP.items():
            if old in source:
                record[new] = source.get(old)
        record.update(_analysis_cost_fields(source))

        rm_source = source.get("QUALITY_rm_source")
        if rm_source == "EVENT_LOG_RECONSTRUCTED":
            for column, value in source.items():
                if column.startswith("RM_") and column.endswith("_t") and column != "RM_total_mass_t":
                    record[f"RM_Reconstructed_{column[3:]}"] = value
        add_source = source.get("QUALITY_add_source")
        if add_source == "EVENT_LOG_RECONSTRUCTED":
            for column, value in source.items():
                if column.startswith("ADD_") and column.endswith("_kg") and column != "ADD_total_kg":
                    record[f"ADD_Reconstructed_{column[4:]}"] = value

        for column, value in source.items():
            match = re.fullmatch(r"PWR_(.+MW)_duration_sec", column)
            if match:
                record[f"Power_{match.group(1)}_sec"] = value
            basket = re.fullmatch(r"BASKET_(\d+)_(time_sec|mass_t)", column)
            if basket and int(basket.group(1)) <= 3:
                suffix = "Time_sec" if basket.group(2) == "time_sec" else "Mass_t"
                record[f"Basket{basket.group(1)}_{suffix}"] = value
            gap = re.fullmatch(r"BASKET_(\d+)_to_(\d+)_gap_sec", column)
            if gap and int(gap.group(1)) <= 2:
                record[f"Basket{gap.group(1)}_to_{gap.group(2)}_Gap_sec"] = value

        _analysis_other_cost_fields(record, source)
        records.append(record)

    frame = pd.DataFrame(records)
    if frame.empty:
        return pd.DataFrame(columns=["Run_ID"])
    frame, _ = drop_all_null_export_columns(frame)
    preferred = ["Run_ID"] + list(ANALYSIS_FIELD_MAP.values()) + [
        "Other_Consumables_Raw_Residual", "Other_Consumables_USD",
        "Cost_Rounding_Adjustment_USD", "Cost_Reconciliation_Error_USD",
        "Other_Cost_Reconciliation_Error_USD",
    ]
    ordered = [column for column in preferred if column in frame]
    remaining = sorted(
        (column for column in frame if column not in ordered),
        key=_natural_column_key,
    )
    return frame.reindex(columns=ordered + remaining)


def build_excel_frames(
    runs_df: pd.DataFrame,
    logs_df: pd.DataFrame,
    ml_features_df: Optional[pd.DataFrame] = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    results = build_results_dataframe(runs_df, logs_df)
    analysis = build_analysis_dataframe(runs_df, logs_df, ml_features_df)
    return results, analysis


def _header_group(column: str, sheet_name: str) -> str:
    if sheet_name == "RESULTS":
        if column.startswith("RM_"):
            return "RM"
        if column.startswith("ADD_"):
            return "ADD"
        if column.startswith("Steel_"):
            return "STEEL"
        if column.startswith("Slag_"):
            return "SLAG"
        if column.startswith("Event_Log_"):
            return "EVENT"
        if column.startswith("Reported_") or column in {"Total_Cost", "Cost_Per_Tonne"}:
            return "COST"
        if column in {
            "Time_min", "Tapping_Mass_t", "Tap_Temperature_C",
            "Total_Energy_kWh", "Energy_kWh_per_t",
        }:
            return "PERF"
        return "RUN"
    if column.startswith(("RM_", "ADD_")):
        return "RECON"
    if column.endswith("_Pass") or column in {
        "Validation_Available", "Unknown_Event_Count", "Warning_Count", "Warning_Text"
    }:
        return "QUALITY"
    if "Cost" in column or "Consumables" in column or column.startswith("Reported_Scrap"):
        return "COST"
    return "PROCESS"


def apply_excel_formatting(
    writer: pd.ExcelWriter, frame: pd.DataFrame, sheet_name: str
) -> None:
    worksheet = writer.book[sheet_name]
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 85
    colors = {
        "RUN": "1F4E78", "PERF": "0F6B5D", "COST": "8A5A00",
        "RM": "355E3B", "ADD": "526D82", "STEEL": "4C566A",
        "SLAG": "5B4B8A", "EVENT": "334155", "RECON": "7C3AED",
        "PROCESS": "0369A1", "QUALITY": "7F1D1D",
    }
    for cell in worksheet[1]:
        cell.fill = PatternFill(
            "solid", fgColor=colors[_header_group(str(cell.value), sheet_name)]
        )
        cell.font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 36
    for row_index in range(2, len(frame) + 2):
        worksheet.row_dimensions[row_index].height = 20

    sample_rows = min(len(frame), 200)
    for index, column in enumerate(frame.columns, start=1):
        letter = get_column_letter(index)
        lengths = [len(str(column))]
        if sample_rows:
            lengths.extend(
                len(str(value)) for value in frame[column].head(sample_rows)
                if not _is_missing(value)
            )
        raw_event = str(column).startswith("Event_Log_Raw_")
        worksheet.column_dimensions[letter].width = (
            48 if raw_event or column == "Warning_Text"
            else min(max(max(lengths, default=8) + 2, 11), 24)
        )
        if raw_event or column == "Warning_Text":
            for cell in worksheet[letter][1:]:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
        if column.endswith("_USD") or column in {
            "Total_Cost", "Cost_Per_Tonne", "Reported_Power", "Reported_Scrap",
            "Reported_Additions", "Reported_Other_Consumables",
        }:
            for cell in worksheet[letter][1:]:
                cell.number_format = '#,##0.00'

    if sheet_name == "ANALYSIS":
        green = PatternFill("solid", fgColor="DCFCE7")
        red = PatternFill("solid", fgColor="FEE2E2")
        for index, column in enumerate(frame.columns, start=1):
            if column.endswith("_Pass"):
                letter = get_column_letter(index)
                target = f"{letter}2:{letter}{max(2, len(frame) + 1)}"
                worksheet.conditional_formatting.add(
                    target, CellIsRule(operator="equal", formula=["1"], fill=green)
                )
                worksheet.conditional_formatting.add(
                    target, CellIsRule(operator="equal", formula=["0"], fill=red)
                )


def make_excel(
    runs_df: pd.DataFrame,
    logs_df: pd.DataFrame,
    ml_features_df: Optional[pd.DataFrame] = None,
) -> BytesIO:
    results, analysis = build_excel_frames(runs_df, logs_df, ml_features_df)
    return make_excel_from_frames(results, analysis)


def make_excel_from_frames(results: pd.DataFrame, analysis: pd.DataFrame) -> BytesIO:
    """Serialize the exact two-sheet public workbook schema."""

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        clean_dataframe_for_excel(results).to_excel(
            writer, sheet_name="RESULTS", index=False
        )
        clean_dataframe_for_excel(analysis).to_excel(
            writer, sheet_name="ANALYSIS", index=False
        )
        apply_excel_formatting(writer, results, "RESULTS")
        apply_excel_formatting(writer, analysis, "ANALYSIS")
    output.seek(0)
    return output


def make_excel_from_frame(all_runs: pd.DataFrame) -> BytesIO:
    """Legacy helper retained for callers that explicitly need an internal dump."""

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        clean_dataframe_for_excel(all_runs).to_excel(
            writer, sheet_name="INTERNAL_FEATURES", index=False
        )
    output.seek(0)
    return output


def summarize_export_frame(frame: pd.DataFrame) -> dict[str, int]:
    def count_equals(column: str, value: str) -> int:
        if column not in frame:
            return 0
        return int(frame[column].fillna("").astype(str).eq(value).sum())

    return {
        "runs": len(frame),
        "columns": len(frame.columns),
        "merged_suspects": int(
            pd.to_numeric(
                frame.get("QUALITY_multi_run_merged_detected", pd.Series(dtype=float)),
                errors="coerce",
            ).fillna(0).sum()
        ),
        "rm_report": count_equals("QUALITY_rm_source", "REPORT")
        + count_equals("QUALITY_rm_source", "REPORT_AND_LOG_VERIFIED"),
        "rm_reconstructed": count_equals(
            "QUALITY_rm_source", "EVENT_LOG_RECONSTRUCTED"
        ),
        "add_report": count_equals("QUALITY_add_source", "REPORT")
        + count_equals("QUALITY_add_source", "REPORT_AND_LOG_VERIFIED"),
        "add_reconstructed": count_equals(
            "QUALITY_add_source", "EVENT_LOG_RECONSTRUCTED"
        ),
        "cost_mapping_warnings": int(
            pd.to_numeric(
                frame.get("COST_mapping_applied", pd.Series(dtype=float)),
                errors="coerce",
            ).fillna(0).sum()
        ),
        "warning_runs": int(
            pd.to_numeric(
                frame.get("QUALITY_warning_count", pd.Series(dtype=float)),
                errors="coerce",
            ).fillna(0).gt(0).sum()
        ),
    }


def summarize_export_frames(
    results: pd.DataFrame, analysis: pd.DataFrame
) -> dict[str, int]:
    """Return UI diagnostics for the public two-sheet workbook."""

    def count_equals(column: str, value: str) -> int:
        if column not in analysis:
            return 0
        return int(analysis[column].fillna("").astype(str).eq(value).sum())

    mapping = pd.to_numeric(
        analysis.get("Cost_Mapping_Applied", pd.Series(dtype=float)),
        errors="coerce",
    ).fillna(0)
    warnings = pd.to_numeric(
        analysis.get("Warning_Count", pd.Series(dtype=float)), errors="coerce"
    ).fillna(0)
    return {
        "runs": len(results),
        "results_columns": len(results.columns),
        "analysis_columns": len(analysis.columns),
        "columns": len(results.columns) + len(analysis.columns),
        "merged_suspects": int(
            pd.to_numeric(
                analysis.get("Multi_Run_Merged_Suspect", pd.Series(dtype=float)),
                errors="coerce",
            ).fillna(0).sum()
        ),
        "rm_report": count_equals("RM_Source", "REPORT")
        + count_equals("RM_Source", "REPORT_AND_LOG_VERIFIED"),
        "rm_reconstructed": count_equals("RM_Source", "EVENT_LOG_RECONSTRUCTED"),
        "add_report": count_equals("ADD_Source", "REPORT")
        + count_equals("ADD_Source", "REPORT_AND_LOG_VERIFIED"),
        "add_reconstructed": count_equals("ADD_Source", "EVENT_LOG_RECONSTRUCTED"),
        "cost_mapping_warnings": int(mapping.sum()),
        "warning_runs": int(warnings.gt(0).sum()),
    }
