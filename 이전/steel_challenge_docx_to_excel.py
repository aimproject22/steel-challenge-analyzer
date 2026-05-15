# app.py
# -*- coding: utf-8 -*-

import re
import json
import sqlite3
from io import BytesIO
from pathlib import Path
from datetime import datetime
from urllib.parse import unquote

import pandas as pd
import streamlit as st
from docx import Document


DB_PATH = Path("steel_challenge.db")


SECTION_RUN = "Run Information"
SECTION_SIM = "Simulation Settings"
SECTION_COST = "Cost Breakdown"
SECTION_STEEL = "Steel Composition"
SECTION_RAW = "Raw Materials"
SECTION_ADD = "Additions"
SECTION_SLAG = "Slag Composition"

HEADER_WORDS = {
    "Element", "Current", "Min", "Max", "Name", "Qty [t]", "Qty[t]"
}


# =========================
# 기본 유틸
# =========================
def normalize_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_value(value):
    text = normalize_text(value)

    if text == "":
        return None

    text = text.replace("$", "")
    text = text.replace(",", "")
    text = text.replace("kWh/t", "")
    text = text.replace("kWh", "")
    text = text.replace("kg", "")
    text = re.sub(r"(?<=\d)t\b", "", text)
    text = text.strip()

    if re.fullmatch(r"[-+]?\d+(\.\d+)?", text):
        num = float(text)
        return int(num) if num.is_integer() else num

    return text


def time_to_minutes(time_text):
    """
    00:18:21 -> 18.35분
    """
    try:
        h, m, s = map(int, time_text.split(":"))
        return h * 60 + m + s / 60
    except Exception:
        return None


# =========================
# DOCX 파싱
# =========================
def table_rows(doc, table_index):
    if table_index >= len(doc.tables):
        return []

    rows = []

    for row in doc.tables[table_index].rows:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)

    return rows


def parse_key_value_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}
    total_energy_count = 0

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_text(row[0])
        value = clean_value(row[1])

        if not key or key in HEADER_WORDS:
            continue

        if section_name == SECTION_COST and key == "Total Energy":
            total_energy_count += 1
            save_key = "Total Energy" if total_energy_count == 1 else "Total Energy_2"
            data[f"{section_name} > {save_key}"] = value
        else:
            data[f"{section_name} > {key}"] = value

    return data


def parse_current_min_max_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_text(row[0])

        if not key or key in HEADER_WORDS:
            continue

        data[f"{section_name} > {key}"] = clean_value(row[1])

    return data


def parse_event_log(doc):
    logs = []
    text_parts = []

    for p in doc.paragraphs:
        if p.text:
            text_parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    text_parts.append(cell.text)

    full_text = "\n".join(text_parts)

    if "Event Log" not in full_text:
        return logs

    event_text = full_text.split("Event Log", 1)[1]

    matches = re.findall(
        r'"(\d{2}:\d{2}:\d{2},.*?)"',
        event_text,
        flags=re.DOTALL
    )

    for idx, item in enumerate(matches, start=1):
        item = item.replace("\n", "")
        item = item.strip().rstrip(",")

        if "," not in item:
            continue

        time_part, event_part = item.split(",", 1)

        logs.append({
            "log_no": idx,
            "time": normalize_text(time_part),
            "event": unquote(normalize_text(event_part)),
        })

    return logs


def parse_docx(uploaded_file):
    doc = Document(uploaded_file)

    if len(doc.tables) < 7:
        raise ValueError("DOCX 내부 표 개수가 부족합니다. Steel Challenge 결과 파일인지 확인하세요.")

    data = {}

    data.update(parse_key_value_table(doc, 0, SECTION_RUN))
    data.update(parse_key_value_table(doc, 1, SECTION_SIM))
    data.update(parse_key_value_table(doc, 2, SECTION_COST))
    data.update(parse_current_min_max_table(doc, 3, SECTION_STEEL))
    data.update(parse_key_value_table(doc, 4, SECTION_RAW))
    data.update(parse_key_value_table(doc, 5, SECTION_ADD))
    data.update(parse_current_min_max_table(doc, 6, SECTION_SLAG))

    logs = parse_event_log(doc)

    return data, logs


# =========================
# ML용 Feature 추출
# =========================
def extract_features_from_logs(logs):
    features = {
        "feature > log_count": len(logs),

        "feature > power_change_count": 0,
        "feature > oxygen_change_count": 0,
        "feature > scrap_basket_count": 0,
        "feature > addition_event_count": 0,
        "feature > analysis_requested_count": 0,
        "feature > analysis_received_count": 0,

        "feature > power_120_first_min": None,
        "feature > power_75_first_min": None,
        "feature > power_0_first_min": None,

        "feature > oxygen_first_min": None,
        "feature > oxygen_150_first_min": None,
        "feature > oxygen_0_first_min": None,

        "feature > tapping_start_min": None,
        "feature > tapping_complete_min": None,
        "feature > tapping_duration_min": None,

        "feature > carbon_added_kg_log": 0,
        "feature > lime_added_kg_log": 0,
        "feature > dolomite_added_kg_log": 0,
        "feature > chrome_carbure_low_s_added_kg_log": 0,
        "feature > silico_chromium_added_kg_log": 0,
    }

    for log in logs:
        time_text = log["time"]
        event = log["event"]
        t_min = time_to_minutes(time_text)

        if "Power set to" in event:
            features["feature > power_change_count"] += 1

            if "120 MW" in event and features["feature > power_120_first_min"] is None:
                features["feature > power_120_first_min"] = t_min

            if "75 MW" in event and features["feature > power_75_first_min"] is None:
                features["feature > power_75_first_min"] = t_min

            if "0 MW" in event and features["feature > power_0_first_min"] is None:
                features["feature > power_0_first_min"] = t_min

        if "Oxygen flow changed" in event:
            features["feature > oxygen_change_count"] += 1

            if features["feature > oxygen_first_min"] is None:
                features["feature > oxygen_first_min"] = t_min

            if "150" in event and features["feature > oxygen_150_first_min"] is None:
                features["feature > oxygen_150_first_min"] = t_min

            if "(0" in event and features["feature > oxygen_0_first_min"] is None:
                features["feature > oxygen_0_first_min"] = t_min

        if "Scrap basket added" in event:
            features["feature > scrap_basket_count"] += 1

        if "Analysis Requested" in event:
            features["feature > analysis_requested_count"] += 1

        if "Analysis received" in event:
            features["feature > analysis_received_count"] += 1

        if "Tapping start" in event:
            features["feature > tapping_start_min"] = t_min

        if "Tapping complete" in event:
            features["feature > tapping_complete_min"] = t_min

        if "Additions:" in event:
            features["feature > addition_event_count"] += 1

            patterns = {
                "feature > carbon_added_kg_log": r"Carbon:\s*(\d+)\s*kg",
                "feature > lime_added_kg_log": r"Lime:\s*(\d+)\s*kg",
                "feature > dolomite_added_kg_log": r"Dolomite:\s*(\d+)\s*kg",
                "feature > chrome_carbure_low_s_added_kg_log": r"Chrome-Carbure Low S:\s*(\d+)\s*kg",
                "feature > silico_chromium_added_kg_log": r"Silico-Chromium:\s*(\d+)\s*kg",
            }

            for feature_name, pattern in patterns.items():
                match = re.search(pattern, event)
                if match:
                    features[feature_name] += int(match.group(1))

    if (
        features["feature > tapping_start_min"] is not None
        and features["feature > tapping_complete_min"] is not None
    ):
        features["feature > tapping_duration_min"] = (
            features["feature > tapping_complete_min"]
            - features["feature > tapping_start_min"]
        )

    return features


# =========================
# SQLite DB
# =========================
def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            uploaded_at TEXT,
            uploader TEXT,
            file_name TEXT,
            score REAL,
            cost_per_tonne REAL,
            steel_grade TEXT,
            data_json TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id INTEGER,
            log_no INTEGER,
            time TEXT,
            event TEXT,
            FOREIGN KEY(run_id) REFERENCES runs(id)
        )
    """)

    conn.commit()
    conn.close()


def save_to_db(uploader, file_name, data, logs):
    features = extract_features_from_logs(logs)

    merged_data = {}
    merged_data.update(data)
    merged_data.update(features)

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    score = data.get("Run Information > Score")
    cost_per_tonne = data.get("Cost Breakdown > Cost Per Tonne")
    steel_grade = data.get("Simulation Settings > Steel Grade")

    cur.execute("""
        INSERT INTO runs (
            uploaded_at,
            uploader,
            file_name,
            score,
            cost_per_tonne,
            steel_grade,
            data_json
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        uploader,
        file_name,
        score,
        cost_per_tonne,
        steel_grade,
        json.dumps(merged_data, ensure_ascii=False)
    ))

    run_id = cur.lastrowid

    for log in logs:
        cur.execute("""
            INSERT INTO logs (
                run_id,
                log_no,
                time,
                event
            )
            VALUES (?, ?, ?, ?)
        """, (
            run_id,
            log["log_no"],
            log["time"],
            log["event"]
        ))

    conn.commit()
    conn.close()


def load_runs_df():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query("SELECT * FROM runs ORDER BY id DESC", conn)
    conn.close()

    if df.empty:
        return df

    rows = []

    for _, row in df.iterrows():
        data = json.loads(row["data_json"])

        new_row = {
            "Run ID": row["id"],
            "Uploaded At": row["uploaded_at"],
            "Uploader": row["uploader"],
            "File Name": row["file_name"],
            "Score": row["score"],
            "Cost Per Tonne": row["cost_per_tonne"],
            "Steel Grade": row["steel_grade"],
        }

        new_row.update(data)
        rows.append(new_row)

    return pd.DataFrame(rows)


def load_logs_df():
    conn = sqlite3.connect(DB_PATH)

    query = """
        SELECT
            runs.id AS run_id,
            runs.uploaded_at,
            runs.uploader,
            runs.file_name,
            logs.log_no,
            logs.time,
            logs.event
        FROM logs
        JOIN runs ON logs.run_id = runs.id
        ORDER BY runs.id DESC, logs.log_no ASC
    """

    df = pd.read_sql_query(query, conn)
    conn.close()

    return df


def make_excel(runs_df, logs_df):
    output = BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        runs_df.to_excel(writer, sheet_name="Results_and_Features", index=False)
        logs_df.to_excel(writer, sheet_name="Event_Log", index=False)

    output.seek(0)
    return output


def reset_database():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("DELETE FROM logs")
    cur.execute("DELETE FROM runs")
    conn.commit()
    conn.close()


# =========================
# Streamlit UI
# =========================
st.set_page_config(
    page_title="Steel Challenge DOCX Analyzer",
    layout="wide"
)

init_db()

st.title("Steel Challenge DOCX 자동 정리 플랫폼")
st.caption("DOCX 업로드 → DB 저장 → 공용 결과 확인 → ML용 Feature 자동 추출")

uploader = st.text_input("업로드한 사람 이름", placeholder="예: 하준영")

uploaded_files = st.file_uploader(
    "Steel Challenge 결과 DOCX 파일 업로드",
    type=["docx"],
    accept_multiple_files=True
)

if st.button("업로드 파일 저장"):
    if not uploader:
        st.warning("업로드한 사람 이름을 입력하세요.")
    elif not uploaded_files:
        st.warning("DOCX 파일을 업로드하세요.")
    else:
        success_count = 0

        for uploaded_file in uploaded_files:
            try:
                data, logs = parse_docx(uploaded_file)
                save_to_db(uploader, uploaded_file.name, data, logs)
                success_count += 1
            except Exception as e:
                st.error(f"{uploaded_file.name} 처리 실패: {e}")

        if success_count > 0:
            st.success(f"{success_count}개 파일 저장 완료")
            st.rerun()


st.divider()

runs_df = load_runs_df()
logs_df = load_logs_df()

st.subheader("전체 업로드 결과 + ML용 Feature")

if runs_df.empty:
    st.info("아직 저장된 결과가 없습니다.")
else:
    feature_cols = [c for c in runs_df.columns if str(c).startswith("feature >")]
    base_cols = [c for c in runs_df.columns if not str(c).startswith("feature >")]

    tab1, tab2, tab3 = st.tabs(["전체 결과", "ML Feature", "Event Log"])

    with tab1:
        st.dataframe(runs_df[base_cols], use_container_width=True)

    with tab2:
        if feature_cols:
            st.dataframe(
                runs_df[["Run ID", "Uploader", "File Name", "Score", "Cost Per Tonne"] + feature_cols],
                use_container_width=True
            )
        else:
            st.info("아직 추출된 feature가 없습니다.")

    with tab3:
        st.dataframe(logs_df, use_container_width=True)

    excel_file = make_excel(runs_df, logs_df)

    st.download_button(
        label="전체 통합 엑셀 다운로드",
        data=excel_file,
        file_name="steel_challenge_all_results_with_features.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


st.divider()

with st.expander("관리자용: 전체 데이터 초기화"):
    password = st.text_input("관리자 비밀번호", type="password")

    if st.button("DB 전체 초기화"):
        if password == "1234":
            reset_database()
            st.success("DB 초기화 완료")
            st.rerun()
        else:
            st.error("비밀번호가 틀렸습니다.")