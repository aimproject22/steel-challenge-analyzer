# -*- coding: utf-8 -*-
"""Steel Challenge DOCX parser.

``parse_docx`` keeps the original ``(data, logs)`` return contract. The
canonical parser additionally preserves Current/Min/Max values and warnings.
"""

from __future__ import annotations

import json
import re
from typing import Any, Iterable, List, Optional
from urllib.parse import unquote

from docx import Document

from canonical import (
    HEADER_WORDS,
    SECTION_ADD,
    SECTION_COST,
    SECTION_RAW,
    SECTION_RUN,
    SECTION_SIM,
    SECTION_SLAG,
    SECTION_STEEL,
    ParsedRun,
    add_measurement,
    build_structured_sections,
    normalize_event_log,
    normalize_key,
    normalize_text,
    normalize_value,
)


# Backward-compatible public name used by older scripts.
clean_value = normalize_value


def table_rows(doc: Any, table_index: int) -> List[List[str]]:
    if table_index >= len(doc.tables):
        return []
    return _rows_from_table(doc.tables[table_index])


def _rows_from_table(table: Any) -> List[List[str]]:
    rows: List[List[str]] = []
    for row in table.rows:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return rows


def _set_duplicate_safe(data: dict, section_name: str, key: str, value: Any) -> None:
    base_key = f"{section_name} > {key}"
    if base_key not in data:
        data[base_key] = value
        return

    suffix = 2
    # Preserve the historical Total Energy_2 spelling.
    candidate = f"{section_name} > {key}_{suffix}"
    while candidate in data:
        suffix += 1
        candidate = f"{section_name} > {key}_{suffix}"
    data[candidate] = value


def parse_key_value_rows(rows: Iterable[Iterable[Any]], section_name: str) -> dict:
    data: dict = {}
    for raw_row in rows:
        row = list(raw_row)
        if len(row) < 2:
            continue
        key = normalize_key(row[0])
        if not key or key.casefold() in HEADER_WORDS:
            continue
        _set_duplicate_safe(data, section_name, key, normalize_value(row[1]))
    return data


def parse_key_value_table(doc: Any, table_index: int, section_name: str) -> dict:
    return parse_key_value_rows(table_rows(doc, table_index), section_name)


def parse_current_min_max_rows(rows: Iterable[Iterable[Any]], section_name: str) -> dict:
    data: dict = {}
    for raw_row in rows:
        row = list(raw_row)
        if len(row) < 2:
            continue
        key = normalize_key(row[0])
        if not key or key.casefold() in HEADER_WORDS:
            continue
        add_measurement(
            data,
            section_name,
            key,
            row[1] if len(row) > 1 else None,
            row[2] if len(row) > 2 else None,
            row[3] if len(row) > 3 else None,
        )
    return data


def parse_current_min_max_table(doc: Any, table_index: int, section_name: str) -> dict:
    return parse_current_min_max_rows(table_rows(doc, table_index), section_name)


def _classify_table(rows: List[List[str]]) -> Optional[str]:
    if not rows:
        return None
    first_col = {normalize_key(row[0]).casefold() for row in rows if row}
    header = [normalize_text(cell).casefold() for cell in rows[0]]

    if {"element", "current"}.issubset(set(header)):
        if {"al2o3", "cao", "feo", "basicity"} & first_col:
            return SECTION_SLAG
        return SECTION_STEEL
    if "user id" in first_col or {"date", "score"}.issubset(first_col):
        return SECTION_RUN
    if "user level" in first_col or "steel grade" in first_col:
        return SECTION_SIM
    if "cost per tonne" in first_col or "tap temperature" in first_col:
        return SECTION_COST
    if header and header[0] == "name":
        addition_markers = (
            "ferro-",
            "lime",
            "dolomite",
            "fluorspar",
            "iron oxide",
            "aluminum pebbles",
            "carbon",
        )
        if any(any(marker in value for marker in addition_markers) for value in first_col):
            return SECTION_ADD
        return SECTION_RAW
    return None


def _extract_event_items(text: str) -> List[str]:
    if "Event Log" not in text:
        return []
    event_text = text.split("Event Log", 1)[1]

    bracket_match = re.search(r"\[(.*)\]", event_text, flags=re.DOTALL)
    if bracket_match:
        candidate = "[" + bracket_match.group(1) + "]"
        try:
            parsed = json.loads(candidate)
            if isinstance(parsed, list):
                return [str(item) for item in parsed]
        except (json.JSONDecodeError, TypeError):
            pass

    return re.findall(
        r'["\'](\d{2}:\d{2}:\d{2},.*?)["\']',
        event_text,
        flags=re.DOTALL,
    )


def parse_event_log(doc: Any) -> list[dict]:
    paragraph_text = "\n".join(p.text for p in doc.paragraphs if p.text)
    full_text = paragraph_text
    if "Event Log" not in full_text:
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        table_text.append(cell.text)
        full_text = "\n".join([paragraph_text, *table_text])

    logs: list[dict] = []
    for index, item in enumerate(_extract_event_items(full_text), start=1):
        item = item.replace("\n", "").strip().rstrip(",")
        if "," not in item:
            continue
        time_part, event_part = item.split(",", 1)
        logs.append(
            normalize_event_log(
                {
                    "log_no": index,
                    "time": normalize_text(time_part),
                    "event": unquote(normalize_text(event_part)),
                },
                index,
            )
        )
    return logs


def parse_docx_canonical(uploaded_file: Any) -> ParsedRun:
    doc = Document(uploaded_file)
    if not doc.tables:
        raise ValueError("DOCX 내부 표가 없습니다. Steel Challenge 결과 파일인지 확인하세요.")

    data: dict = {}
    warnings: list[str] = []
    seen_sections: set[str] = set()

    for table in doc.tables:
        rows = _rows_from_table(table)
        section = _classify_table(rows)
        if not section:
            warnings.append("인식하지 못한 표를 건너뛰었습니다.")
            continue
        if section in seen_sections:
            warnings.append(f"중복 섹션을 병합했습니다: {section}")
        seen_sections.add(section)
        if section in {SECTION_STEEL, SECTION_SLAG}:
            data.update(parse_current_min_max_rows(rows, section))
        else:
            parsed = parse_key_value_rows(rows, section)
            for key, value in parsed.items():
                if key in data:
                    section_name, item_name = key.split(" > ", 1)
                    _set_duplicate_safe(data, section_name, item_name, value)
                else:
                    data[key] = value

    # Raw Materials and Additions may be absent in otherwise valid EAF reports;
    # the export layer can reconstruct both from the Event Log.
    required = {SECTION_RUN, SECTION_SIM, SECTION_COST, SECTION_STEEL, SECTION_SLAG}
    missing = sorted(required - seen_sections)
    if missing:
        raise ValueError(f"필수 섹션이 없습니다: {', '.join(missing)}")
    if SECTION_RAW not in seen_sections:
        warnings.append("Raw Materials 섹션이 없습니다.")

    for section, label in ((SECTION_RAW, "Raw Materials"), (SECTION_ADD, "Additions")):
        parsed = any(str(key).startswith(f"{section} > ") for key in data)
        data[f"Parser Diagnostics > {label} Parse Status"] = (
            "REPORT_PARSED"
            if parsed
            else "REPORT_PARSE_FAILED"
            if section in seen_sections
            else "REPORT_SECTION_MISSING"
        )

    data.update(build_structured_sections(data))
    logs = parse_event_log(doc)
    if not logs:
        warnings.append("Event Log를 찾지 못했습니다.")
    return ParsedRun(data=data, logs=logs, warnings=warnings)


def parse_docx(uploaded_file: Any):
    """Return the legacy ``(data, logs)`` tuple used by the existing app."""

    return parse_docx_canonical(uploaded_file).to_legacy_tuple()
