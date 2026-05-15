# parser.py
# -*- coding: utf-8 -*-

import re
from urllib.parse import unquote
from docx import Document


SECTION_RUN = "Run Information"
SECTION_SIM = "Simulation Settings"
SECTION_COST = "Cost Breakdown"
SECTION_STEEL = "Steel Composition"
SECTION_RAW = "Raw Materials"
SECTION_ADD = "Additions"
SECTION_SLAG = "Slag Composition"

HEADER_WORDS = {
    "Element", "Current", "Min", "Max", "Name", "Qty [t]", "Qty[t]"
}


def normalize_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_value(value):
    text = normalize_text(value)

    if text == "":
        return None

    text = text.replace("$", "")
    text = text.replace(",", "")
    text = text.replace("kWh/t", "")
    text = text.replace("kWh", "")
    text = text.replace("kg", "")
    text = re.sub(r"(?<=\d)t\b", "", text)
    text = text.strip()

    if re.fullmatch(r"[-+]?\d+(\.\d+)?", text):
        num = float(text)
        return int(num) if num.is_integer() else num

    return text


def table_rows(doc, table_index):
    if table_index >= len(doc.tables):
        return []

    rows = []

    for row in doc.tables[table_index].rows:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(cells)

    return rows


def parse_key_value_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}
    total_energy_count = 0

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_text(row[0])
        value = clean_value(row[1])

        if not key or key in HEADER_WORDS:
            continue

        if section_name == SECTION_COST and key == "Total Energy":
            total_energy_count += 1
            save_key = "Total Energy" if total_energy_count == 1 else "Total Energy_2"
            data[f"{section_name} > {save_key}"] = value
        else:
            data[f"{section_name} > {key}"] = value

    return data


def parse_current_min_max_table(doc, table_index, section_name):
    rows = table_rows(doc, table_index)
    data = {}

    for row in rows:
        if len(row) < 2:
            continue

        key = normalize_text(row[0])

        if not key or key in HEADER_WORDS:
            continue

        data[f"{section_name} > {key}"] = clean_value(row[1])

    return data


def parse_event_log(doc):
    logs = []
    text_parts = []

    for p in doc.paragraphs:
        if p.text:
            text_parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    text_parts.append(cell.text)

    full_text = "\n".join(text_parts)

    if "Event Log" not in full_text:
        return logs

    event_text = full_text.split("Event Log", 1)[1]

    matches = re.findall(
        r'"(\d{2}:\d{2}:\d{2},.*?)"',
        event_text,
        flags=re.DOTALL
    )

    for idx, item in enumerate(matches, start=1):
        item = item.replace("\n", "")
        item = item.strip().rstrip(",")

        if "," not in item:
            continue

        time_part, event_part = item.split(",", 1)

        logs.append({
            "log_no": idx,
            "time": normalize_text(time_part),
            "event": unquote(normalize_text(event_part)),
        })

    return logs


def parse_docx(uploaded_file):
    doc = Document(uploaded_file)

    if len(doc.tables) < 7:
        raise ValueError("DOCX 내부 표 개수가 부족합니다. Steel Challenge 결과 파일인지 확인하세요.")

    data = {}

    data.update(parse_key_value_table(doc, 0, SECTION_RUN))
    data.update(parse_key_value_table(doc, 1, SECTION_SIM))
    data.update(parse_key_value_table(doc, 2, SECTION_COST))
    data.update(parse_current_min_max_table(doc, 3, SECTION_STEEL))
    data.update(parse_key_value_table(doc, 4, SECTION_RAW))
    data.update(parse_key_value_table(doc, 5, SECTION_ADD))
    data.update(parse_current_min_max_table(doc, 6, SECTION_SLAG))

    logs = parse_event_log(doc)

    return data, logs