import re
from pathlib import Path
from docx import Document
from openpyxl import load_workbook


DOCX_FOLDER = Path("docx_results")          # DOCX 파일 넣을 폴더
TEMPLATE_XLSX = Path("스틸 챌린지 컬럼.xlsx")
OUTPUT_XLSX = Path("스틸_챌린지_정리결과.xlsx")


def read_docx_text(path):
    doc = Document(path)
    lines = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                lines.extend(row_text)

    return lines


def clean_number(value):
    if value is None:
        return None

    value = str(value)
    value = value.replace("$", "").replace(",", "")
    value = value.replace("kWh/t", "").replace("kWh", "")
    value = value.replace("t", "").replace("kg", "")
    value = value.strip()

    try:
        return float(value)
    except:
        return value


def find_next_value(lines, key):
    for i, line in enumerate(lines):
        if line.strip() == key:
            if i + 1 < len(lines):
                return clean_number(lines[i + 1])
    return None


def parse_section_table(lines, section_name, stop_sections):
    result = {}

    try:
        start = lines.index(section_name)
    except ValueError:
        return result

    end = len(lines)
    for stop in stop_sections:
        if stop in lines[start + 1:]:
            idx = lines.index(stop, start + 1)
            end = min(end, idx)

    section = lines[start:end]

    skip_words = {"Element", "Current", "Min", "Max", "Name", "Qty [t]"}

    i = 1
    while i < len(section) - 1:
        key = section[i].strip()

        if key in skip_words or key == "":
            i += 1
            continue

        value = section[i + 1].strip()

        if value in skip_words:
            i += 1
            continue

        result[key] = clean_number(value)
        i += 2

    return result


def parse_docx(path):
    lines = read_docx_text(path)

    data = {}

    # Run Information
    data["status"] = find_next_value(lines, "Status")
    data["Score"] = find_next_value(lines, "Score")

    # Simulation Settings
    data["User Level"] = find_next_value(lines, "User Level")
    data["Steel Grade"] = find_next_value(lines, "Steel Grade")

    # Cost Breakdown
    cost_keys = [
        "Time (in minutes)",
        "Tapping mass",
        "Tap temperature",
        "Total Energy",
        "Power",
        "Scrap",
        "Additions",
        "Other consumables",
        "Total Cost",
        "Cost Per Tonne",
    ]

    total_energy_count = 0
    for i, line in enumerate(lines):
        if line == "Total Energy":
            total_energy_count += 1
            if i + 1 < len(lines):
                if total_energy_count == 1:
                    data["Total Energy"] = clean_number(lines[i + 1])
                elif total_energy_count == 2:
                    data["Total Energy_2"] = clean_number(lines[i + 1])

    for key in cost_keys:
        if key != "Total Energy":
            data[key] = find_next_value(lines, key)

    # Steel Composition
    steel_comp = parse_section_table(
        lines,
        "Steel Composition / wt%",
        ["Raw Materials", "Additions", "Slag Composition", "Event Log"]
    )
    for k, v in steel_comp.items():
        data[k] = v

    # Raw Materials
    raw_materials = parse_section_table(
        lines,
        "Raw Materials",
        ["Additions", "Slag Composition", "Event Log"]
    )
    for k, v in raw_materials.items():
        data[k] = v

    # Additions
    additions = parse_section_table(
        lines,
        "Additions",
        ["Slag Composition", "Event Log"]
    )
    for k, v in additions.items():
        data[k] = v

    # Slag Composition
    slag = parse_section_table(
        lines,
        "Slag Composition",
        ["Event Log"]
    )
    for k, v in slag.items():
        data[k] = v

    return data


def normalize_key(key):
    if key is None:
        return ""

    key = str(key).strip()

    replacements = {
        "Time(in minutes)": "Time (in minutes)",
        "Plate and structural": "Plate and Structural",
        "Plate and structural ": "Plate and Structural",
        "Intemal Low Alloyed": "Internal Low Alloyed",
        "eafScrap Type05": "eafScrapType05",
        "eafScrap Type10": "eafScrapType10",
        "Turnnings": "Turnings",
    }

    return replacements.get(key, key)


def write_to_excel(all_results):
    wb = load_workbook(TEMPLATE_XLSX)
    ws = wb.active

    # 1행: 실행 번호
    for col_idx, file_name in enumerate(all_results.keys(), start=2):
        ws.cell(row=1, column=col_idx).value = file_name

    # A열 항목명 기준으로 값 입력
    for row in range(2, ws.max_row + 1):
        raw_key = ws.cell(row=row, column=1).value
        key = normalize_key(raw_key)

        if not key:
            continue

        for col_idx, result in enumerate(all_results.values(), start=2):
            ws.cell(row=row, column=col_idx).value = result.get(key)

    wb.save(OUTPUT_XLSX)


def main():
    docx_files = sorted(DOCX_FOLDER.glob("*.docx"))

    if not docx_files:
        print("DOCX 파일이 없습니다. docx_results 폴더에 파일을 넣어주세요.")
        return

    all_results = {}

    for path in docx_files:
        print(f"처리 중: {path.name}")
        all_results[path.stem] = parse_docx(path)

    write_to_excel(all_results)
    print(f"완료: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()