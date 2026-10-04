from __future__ import annotations

from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_utils import (
    apply_other_consumables_decomposition,
    build_analysis_dataframe,
    build_all_runs_export_df,
    build_excel_frames,
    build_results_dataframe,
    make_excel,
)
from export_features import extract_event_export_features
from ml_engine import get_safe_feature_columns, train_model


def test_ml_minimum_rows_and_target_leakage() -> None:
    small = pd.DataFrame(
        {"Cost Per Tonne": range(10), "feature > useful": range(10)}
    )
    model, _, _, metrics = train_model(small, "Cost Per Tonne")
    assert model is None
    assert "20" in metrics["error"]

    rng = np.random.default_rng(42)
    useful = np.arange(30, dtype=float)
    frame = pd.DataFrame(
        {
            "Run ID": range(30),
            "Cost Per Tonne": 400 - useful + rng.normal(0, 0.2, 30),
            "feature > useful": useful,
            "feature > total_cost_proxy": 400 - useful,
            "feature > sparse": [None if i % 3 == 0 else i for i in range(30)],
        }
    )
    model, features, importance, metrics = train_model(frame, "Cost Per Tonne")
    assert model is not None
    assert "feature > useful" in features
    assert "feature > total_cost_proxy" not in features
    assert set(("R2", "MAE", "RMSE")).issubset(metrics)
    assert not importance.empty


def _headers(worksheet) -> dict[str, int]:
    return {
        str(cell.value): index
        for index, cell in enumerate(worksheet[1], start=1)
        if cell.value is not None
    }


def test_excel_has_exact_two_sheets_and_sanitizes_strings() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Sender": ["=HYPERLINK(\"bad\")"],
            "feature > event_count": [2],
        }
    )
    logs = pd.DataFrame(
        {"run_id": [1], "log_no": [1], "time": ["00:00:00"], "event": ["Power\x00 on"]}
    )
    output = make_excel(runs, logs)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=False)
    assert workbook.sheetnames == ["RESULTS", "ANALYSIS"]
    worksheet = workbook["RESULTS"]
    assert worksheet.freeze_panes == "B2"
    assert worksheet.auto_filter.ref == worksheet.dimensions
    headers = _headers(worksheet)
    assert worksheet.max_row == 2
    assert worksheet.cell(2, headers["Sender"]).value.startswith("'")
    assert "\x00" not in worksheet.cell(2, headers["Event_Log_Raw_01"]).value
    assert worksheet.row_dimensions[2].height == 20
    assert not worksheet.cell(2, headers["Event_Log_Raw_01"]).alignment.wrap_text


def test_excel_two_runs_produce_two_rows_and_aggregate_logs() -> None:
    runs = pd.DataFrame({"Run ID": [1, 2], "Status": [1, 0]})
    logs = pd.DataFrame(
        {
            "run_id": [1, 2],
            "log_no": [1, 1],
            "time": ["00:00:00", "00:00:10"],
            "event": ["Power set to: 120 MW", "Tapping complete"],
        }
    )
    output = make_excel(runs, logs)
    workbook = load_workbook(BytesIO(output.getvalue()), read_only=True)
    worksheet = workbook["RESULTS"]
    headers = _headers(worksheet)
    assert worksheet.max_row == 3
    assert "Power set to" in worksheet.cell(2, headers["Event_Log_Raw_01"]).value
    assert "Tapping complete" in worksheet.cell(3, headers["Event_Log_Raw_01"]).value


def test_dynamic_event_features_power_basket_addition_and_analysis() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Time": [30],
            "Total Energy kWh": [1000],
            "Raw Materials > No1 Bundles": [48],
            "Additions > Lime": [1430],
        }
    )
    event_rows = [
        ("00:00:00", "Power 120 MW"),
        ("00:00:32", "Scrap basket added with: No1 Bundles: 48t"),
        ("00:10:00", "Power 90 MW"),
        ("00:12:00", "Analysis Requested"),
        ("00:13:30", "Analysis received"),
        ("00:15:00", "Additions: Dolomite:930kg; Iron Oxide:700kg; Lime:1430kg"),
        ("00:20:00", "Power 0 MW"),
        ("00:25:00", "Power 180 MW"),
        ("00:30:00", "Tapping complete"),
    ]
    logs = pd.DataFrame(
        [
            {"run_id": 1, "log_no": index, "time": time, "event": event}
            for index, (time, event) in enumerate(event_rows, start=1)
        ]
    )
    frame = build_all_runs_export_df(runs, logs)
    row = frame.iloc[0]
    assert row["PWR_120MW_duration_sec"] == 600
    assert row["PWR_90MW_duration_sec"] == 600
    assert row["PWR_0MW_duration_sec"] == 300
    assert row["PWR_180MW_duration_sec"] == 300
    assert row["PWR_total_on_sec"] == 1500
    assert row["BASKET_1_mass_t"] == 48
    assert row["EVTADD_Dolomite_first_sec"] == 900
    assert row["EVTADD_Dolomite_total_kg"] == 930
    assert row["ANALYSIS_avg_wait_sec"] == 90
    assert row["RM_No1_Bundles_t"] == 48
    assert row["ADD_Lime_kg"] == 1430


def test_power_duration_reference_sequence() -> None:
    events = [
        {"time": "00:00:00", "event": "Power 120 MW"},
        {"time": "00:10:00", "event": "Power 90 MW"},
        {"time": "00:20:00", "event": "Power 0 MW"},
        {"time": "00:30:00", "event": "Tapping complete"},
    ]
    result = extract_event_export_features(events)
    assert result["PWR_120MW_duration_sec"] == 600
    assert result["PWR_90MW_duration_sec"] == 600
    assert result["PWR_0MW_duration_sec"] == 600
    assert result["PWR_total_on_sec"] == 1200


def test_oxygen_and_carbon_flow_integration() -> None:
    events = [
        {"time": "00:00:00", "event": "Oxygen flow changed: (100 Nm³ / min)"},
        {"time": "00:10:00", "event": "Oxygen flow changed: (150 Nm³ / min)"},
        {"time": "00:20:00", "event": "Oxygen flow changed: (0 Nm³ / min)"},
        {"time": "00:30:00", "event": "Carbon flow changed: 2 kg/min"},
        {"time": "00:40:00", "event": "Carbon flow changed: 0 kg/min"},
        {"time": "00:50:00", "event": "Tapping complete"},
    ]
    result = extract_event_export_features(events)
    assert result["O2_integrated_Nm3"] == 2500
    assert result["O2_total_on_sec"] == 1200
    assert result["CINJ_integrated_kg"] == 20
    assert result["CINJ_total_on_sec"] == 600


def test_raw_event_log_splits_without_data_loss() -> None:
    event = "X" * 65_000
    result = extract_event_export_features(
        [{"time": "00:00:00", "event": event}]
    )
    chunks = [
        result[key]
        for key in sorted(result)
        if key.startswith("RAW_event_log_")
    ]
    assert len(chunks) == 3
    assert "".join(chunks) == f"00:00:00,{event}"
    assert all(len(chunk) <= 30_000 for chunk in chunks)


def test_duplicate_parser_columns_are_collapsed_deterministically() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Additions > Lime": [100],
            "Additions > Lime_2": [100],
            "Additions > Lime_3": [200],
        }
    )
    frame = build_all_runs_export_df(runs, pd.DataFrame())
    assert frame.loc[0, "ADD_Lime_kg"] == 100
    assert not any(column.startswith("ADD_Lime_2") for column in frame.columns)
    assert frame.loc[0, "QUALITY_warning_count"] >= 2
    assert "Conflicting duplicate" in frame.loc[0, "QUALITY_warning_text"]


def test_excel_flags_a_database_row_containing_multiple_runs() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1035],
            "Run Information > Date": ["21/09/2026 10:00:00"],
            "Run Information > Date_2": ["21/09/2026 11:00:00"],
            "Run Information > Score": [410.0],
            "Run Information > Score_2": [405.0],
        }
    )

    frame = build_all_runs_export_df(runs, pd.DataFrame())
    assert frame.loc[0, "QUALITY_multi_run_merged_detected"] == 1
    assert "Multiple simulation runs" in frame.loc[0, "QUALITY_warning_text"]


def test_excel_flags_clock_reset_after_tapping_complete() -> None:
    runs = pd.DataFrame({"Run ID": [7]})
    logs = pd.DataFrame(
        {
            "run_id": [7, 7],
            "log_no": [1, 2],
            "time": ["00:40:00", "00:00:05"],
            "event": ["Tapping complete", "Simulation rate changed: 1"],
        }
    )
    frame = build_all_runs_export_df(runs, logs)
    assert frame.loc[0, "QUALITY_multi_run_merged_detected"] == 1


def test_all_null_and_spec_less_columns_are_not_exported() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Slag Composition > Al2O3 > Current": [4.8],
        }
    )
    frame = build_all_runs_export_df(runs, pd.DataFrame())
    assert "SLAG_Al2O3_pct" in frame
    assert "SLAG_Al2O3_min" not in frame
    assert "COST_power_usd" not in frame
    assert not any(frame[column].isna().all() for column in frame)


def test_duplicate_basket_materials_sum_and_reconstruct_report() -> None:
    runs = pd.DataFrame({"Run ID": [1]})
    logs = pd.DataFrame(
        {
            "run_id": [1, 1],
            "log_no": [1, 2],
            "time": ["00:00:00", "00:30:00"],
            "event": [
                "Scrap basket added with: Shredded: 22t; Direct Reduced Iron: 18t; Plate and Structural: 4t; Shredded: 18t",
                "Tapping complete",
            ],
        }
    )
    row = build_all_runs_export_df(runs, logs).iloc[0]
    assert row["BASKET_1_Shredded_t"] == 40
    assert row["BASKET_Shredded_total_t"] == 40
    assert row["RM_Shredded_t"] == 40
    assert row["QUALITY_rm_source"] == "EVENT_LOG_RECONSTRUCTED"


def test_additions_fallback_and_alias_normalization() -> None:
    runs = pd.DataFrame({"Run ID": [1]})
    logs = pd.DataFrame(
        {
            "run_id": [1, 1],
            "log_no": [1, 2],
            "time": ["00:10:00", "00:20:00"],
            "event": [
                "Additions: HC FeMn 600kg; Dolomite 930kg; IronOxide 700kg; Lime 1430kg",
                "Tapping complete",
            ],
        }
    )
    row = build_all_runs_export_df(runs, logs).iloc[0]
    assert row["ADD_High_C_Ferro_Manganese_kg"] == 600
    assert row["ADD_Dolomite_kg"] == 930
    assert row["ADD_Iron_Oxide_kg"] == 700
    assert row["ADD_Lime_kg"] == 1430
    assert row["QUALITY_add_source"] == "EVENT_LOG_RECONSTRUCTED"


def test_cost_mapping_detects_shift_without_price_hardcoding() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Process Type": ["Electric Arc Furnace"],
            "Time": [42],
            "Cost Breakdown > Power": [591.2],
            "Cost Breakdown > Scrap": [42],
            "Cost Breakdown > Additions": [16020],
            "Cost Breakdown > Other consumables": [19300.85],
            "Total Cost": [35912],
        }
    )
    row = build_all_runs_export_df(runs, pd.DataFrame()).iloc[0]
    assert row["COST_mapping_applied"] == 1
    assert row["COST_mapping_type"] == "STEELUNIVERSITY_EAF_KNOWN_SHIFT"
    assert row["COST_source_scrap_is_time_duplicate"] == 1
    assert row["COST_power_usd"] == pytest.approx(19300.85)
    assert row["COST_scrap_usd"] == pytest.approx(16020)
    assert row["COST_additions_usd"] == pytest.approx(591.2)
    assert row["COST_other_consumables_reconstructed_usd"] == pytest.approx(0)
    assert row["COST_reconciliation_error_usd"] == pytest.approx(0.05)


def test_run_1021_cost_mapping_and_electrode_downtime() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1021],
            "Process Type": ["Electric Arc Furnace"],
            "Time": [70],
            "Cost Breakdown > Power": [354.5],
            "Cost Breakdown > Scrap": [70],
            "Cost Breakdown > Additions": [16150],
            "Cost Breakdown > Other consumables": [18968.24],
            "Total Cost": [35771],
        }
    )
    event_rows = [
        ("00:34:54", "Electrode breakage"),
        ("00:46:22", "Broken electrode has been replaced"),
        ("00:56:23", "Oxygen flow changed (10 Nm³ / min)"),
        ("00:56:25", "Oxygen flow changed (20 Nm³ / min)"),
        ("00:56:27", "Oxygen flow changed (30 Nm³ / min)"),
        ("00:56:29", "Oxygen flow changed (40 Nm³ / min)"),
        ("00:56:31", "Oxygen flow changed (50 Nm³ / min)"),
        ("00:56:33", "Oxygen flow changed (60 Nm³ / min)"),
        ("00:56:35", "Oxygen flow changed (70 Nm³ / min)"),
        ("00:56:37", "Oxygen flow changed (80 Nm³ / min)"),
        ("00:56:39", "Oxygen flow changed (90 Nm³ / min)"),
        ("00:56:45", "Oxygen flow changed (100 Nm³ / min)"),
        ("00:56:49", "Oxygen flow changed (110 Nm³ / min)"),
        ("00:56:55", "Oxygen flow changed (120 Nm³ / min)"),
        ("00:57:00", "Oxygen flow changed (130 Nm³ / min)"),
        ("01:02:09", "Oxygen flow changed (120 Nm³ / min)"),
        ("01:02:11", "Oxygen flow changed (110 Nm³ / min)"),
        ("01:02:13", "Oxygen flow changed (100 Nm³ / min)"),
        ("01:02:18", "Oxygen flow changed (90 Nm³ / min)"),
        ("01:06:37", "Oxygen flow changed (0 Nm³ / min)"),
        ("01:08:31", "Tapping start"),
        ("01:10:20", "Tapping complete"),
    ]
    logs = pd.DataFrame(
        [
            {"run_id": 1021, "log_no": index, "time": time, "event": event}
            for index, (time, event) in enumerate(event_rows, start=1)
        ]
    )
    row = build_all_runs_export_df(runs, logs).iloc[0]
    assert row["COST_other_consumables_reconstructed_usd"] == pytest.approx(298.26)
    assert row["ELEC_break_count"] == 1
    assert row["ELEC_first_break_sec"] == 2094
    assert row["ELEC_first_replacement_sec"] == 2782
    assert row["ELEC_total_downtime_sec"] == 688
    assert row["O2_integrated_Nm3"] == pytest.approx(1122.6666666667)
    assert row["CINJ_integrated_kg"] == 0
    assert row["TAP_duration_sec"] == 109


def test_other_consumables_model_uses_data_not_hardcoded_prices() -> None:
    rows = []
    for index in range(20):
        oxygen = 100 + index * 20
        carbon = float((index % 4) * 5)
        rows.append(
            {
                "META_process_type": "Electric Arc Furnace",
                "COST_other_total_usd": 50 + 0.1 * oxygen + 2.0 * carbon,
                "O2_integrated_Nm3": oxygen,
                "O2_used": 1,
                "CINJ_integrated_kg": carbon,
                "CINJ_used": int(carbon > 0),
                "ELEC_break_count": 0,
            }
        )
    rows.append(
        {
            "META_process_type": "Electric Arc Furnace",
            "COST_other_total_usd": 50 + 0.1 * 600 + 2.0 * 10 + 300,
            "O2_integrated_Nm3": 600,
            "O2_used": 1,
            "CINJ_integrated_kg": 10,
            "CINJ_used": 1,
            "ELEC_break_count": 1,
        }
    )
    result = apply_other_consumables_decomposition(pd.DataFrame(rows))
    break_row = result.iloc[-1]
    assert break_row["COST_other_model_name"] == "BASE_O2_C"
    assert break_row["COST_tapping_estimated_usd"] == pytest.approx(50)
    assert break_row["COST_oxygen_estimated_usd"] == pytest.approx(60)
    assert break_row["COST_carbon_injection_estimated_usd"] == pytest.approx(20)
    assert break_row["COST_electrode_breakage_estimated_usd"] == pytest.approx(300)
    assert break_row["COST_other_unexplained_usd"] == pytest.approx(0)


def test_multiple_electrode_events_use_fifo_pairing() -> None:
    result = extract_event_export_features(
        [
            {"time": "00:01:00", "event": "Electrode broke"},
            {"time": "00:02:00", "event": "Electrode breakage"},
            {"time": "00:04:00", "event": "Broken electrode has been replaced"},
            {"time": "00:07:00", "event": "New electrode installed"},
            {"time": "00:08:00", "event": "Tapping complete"},
        ]
    )
    assert result["ELEC_break_count"] == 2
    assert result["ELEC_replacement_count"] == 2
    assert result["ELEC_total_downtime_sec"] == 480
    assert result["ELEC_unpaired_break_count"] == 0
    assert result["ELEC_unpaired_replacement_count"] == 0


def test_non_eaf_costs_are_not_shifted() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [2],
            "Process Type": ["Secondary Steelmaking"],
            "Cost Breakdown > Power": [10],
            "Cost Breakdown > Scrap": [20],
            "Cost Breakdown > Additions": [30],
            "Cost Breakdown > Other consumables": [40],
            "Total Cost": [100],
        }
    )
    row = build_all_runs_export_df(runs, pd.DataFrame()).iloc[0]
    assert row["COST_mapping_applied"] == 0
    assert row["COST_power_usd"] == 10
    assert row["COST_scrap_usd"] == 20
    assert row["COST_additions_usd"] == 30
    assert row["COST_other_consumables_reconstructed_usd"] == 40


def test_chemistry_pass_is_not_mislabeled_as_overall_pass() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Status": [0],
            "Steel Composition > C > Current": [0.11],
            "Steel Composition > C > Min": [0.10],
            "Steel Composition > C > Max": [0.12],
        }
    )
    frame = build_all_runs_export_df(runs, pd.DataFrame())
    assert frame.loc[0, "QUALITY_chemistry_pass"] == 1
    assert frame.loc[0, "QUALITY_full_validation_available"] == 0
    assert "QUALITY_overall_pass" not in frame.columns


def test_power_integral_names_do_not_claim_actual_energy() -> None:
    result = extract_event_export_features(
        [
            {"time": "00:00:00", "event": "Power 120 MW"},
            {"time": "00:10:00", "event": "Tapping complete"},
        ]
    )
    assert result["PWR_setpoint_integral_MW_sec"] == 72_000
    assert result["PWR_setpoint_equivalent_MWh"] == 20
    assert "PWR_energy_estimated_kwh" not in result


def test_ml_roles_exclude_post_run_quality_features() -> None:
    frame = pd.DataFrame(
        {
            "feature > basket_total": [1],
            "feature > steel_c_within_spec": [1],
            "feature > raw_event_sequence": [1],
            "feature > total_cost_proxy": [1],
        }
    )
    assert get_safe_feature_columns(frame, "Cost Per Tonne") == [
        "feature > basket_total"
    ]


def test_excel_formats_both_public_sheets() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Steel Composition > C > Current": [0.11],
            "Steel Composition > C > Min": [0.10],
            "Steel Composition > C > Max": [0.12],
        }
    )
    logs = pd.DataFrame(
        {
            "run_id": [1],
            "log_no": [1],
            "time": ["00:10:00"],
            "event": ["Tapping complete"],
        }
    )
    workbook = load_workbook(BytesIO(make_excel(runs, logs).getvalue()))
    for sheet_name in ("RESULTS", "ANALYSIS"):
        worksheet = workbook[sheet_name]
        assert worksheet.freeze_panes == "B2"
        assert worksheet.auto_filter.ref == worksheet.dimensions
        assert worksheet.sheet_view.showGridLines is False
    headers = _headers(workbook["RESULTS"])
    assert {"Steel_C", "Steel_C_Min", "Steel_C_Max"}.issubset(headers)


def test_reference_success_run_is_preserved() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1029],
            "Status": [1],
            "Score": [406.62],
            "Time": [42],
            "Tapping Mass": [88],
            "Tap Temperature": [1634],
            "Total Energy kWh": [33861],
            "Energy kWh/t": [383],
            "Cost Per Tonne": [406.62],
            "Raw Materials > No1 Bundles": [89],
            "Additions > High C Ferro-Manganese": [600],
            "Additions > Dolomite": [930],
            "Additions > Iron Oxide": [700],
            "Additions > Lime": [1430],
            "Steel Composition > C > Current": [0.115],
            "Steel Composition > C > Min": [0.100],
            "Steel Composition > C > Max": [0.120],
            "Slag Composition > CaO > Current": [43.934],
            "Slag Composition > SiO2 > Current": [25.750],
            "Slag Composition > Basicity > Current": [1.70620239624785],
            "Slag Composition > Basicity > Min": [1.5],
            "Slag Composition > Basicity > Max": [2.5],
        }
    )
    events = [
        ("00:00:32", "Scrap basket added with: No1 Bundles: 48t"),
        ("00:01:35", "Power set to: 200 MW"),
        ("00:01:58", "Power set to: 120 MW"),
        ("00:18:00", "Scrap basket added with: No1 Bundles: 41t"),
        ("00:31:16", "Power set to: 105 MW"),
        ("00:32:15", "Additions: High C Ferro-Manganese: 600kg"),
        ("00:32:51", "Additions: Dolomite: 930kg; Iron Oxide: 700kg; Lime: 1430kg"),
        ("00:40:58", "Tapping start"),
        ("00:42:51", "Tapping complete"),
    ]
    logs = pd.DataFrame(
        [
            {"run_id": 1029, "log_no": index, "time": time, "event": event}
            for index, (time, event) in enumerate(events, start=1)
        ]
    )
    row = build_all_runs_export_df(runs, logs).iloc[0]
    assert row["COST_per_tonne_usd"] == pytest.approx(406.62)
    assert row["PERF_time_min"] == 42
    assert row["PERF_tapping_mass_t"] == 88
    assert row["PERF_tap_temperature_c"] == 1634
    assert row["PERF_total_energy_kwh"] == 33861
    assert row["RM_No1_Bundles_t"] == 89
    assert row["BASKET_No1_Bundles_total_t"] == 89
    assert row["ADD_High_C_Ferro_Manganese_kg"] == 600
    assert row["ADD_Lime_kg"] == 1430
    assert row["STEEL_C_wt_pct"] == pytest.approx(0.115)
    assert row["SLAG_Basicity"] == pytest.approx(1.70620239624785)
    assert row["QUALITY_rm_source"] == "REPORT_AND_LOG_VERIFIED"
    assert row["QUALITY_add_source"] == "REPORT_AND_LOG_VERIFIED"


def test_public_results_contains_reported_values_only() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Time": [42],
            "Cost Breakdown > Power": [591.2],
            "Cost Breakdown > Scrap": [42],
            "Cost Breakdown > Additions": [16020],
            "Cost Breakdown > Other consumables": [19300.85],
            "Raw Materials > No1 Bundles": [89],
            "Additions > Lime": [1430],
            "Steel Composition > C > Current": [0.115],
            "Steel Composition > C > Min": [0.1],
            "Steel Composition > C > Max": [0.12],
        }
    )
    logs = pd.DataFrame(
        {
            "run_id": [1],
            "log_no": [1],
            "time": ["00:10:00"],
            "event": ["Power 120 MW"],
        }
    )
    results = build_results_dataframe(runs, logs)
    assert results.loc[0, "Reported_Power"] == pytest.approx(591.2)
    assert results.loc[0, "Reported_Scrap"] == pytest.approx(42)
    assert results.loc[0, "RM_No1_Bundles_t"] == 89
    assert results.loc[0, "ADD_Lime_kg"] == 1430
    assert results.loc[0, "Steel_C"] == pytest.approx(0.115)
    forbidden = (
        "QUALITY_", "BASKET_", "PWR_", "O2_", "CINJ_", "ELEC_",
        "EVTADD_", "ANALYSIS_", "TAP_", "DERIVED_",
    )
    assert not any(column.startswith(forbidden) for column in results.columns)
    assert not any("Corrected" in column or "Reconstructed" in column for column in results)


def test_two_sheets_keep_identical_run_order_and_one_row_per_run() -> None:
    runs = pd.DataFrame({"Run ID": [1003, 1001, 1002], "Status": [1, 0, 1]})
    results, analysis = build_excel_frames(runs, pd.DataFrame())
    assert results["Run_ID"].tolist() == [1003, 1001, 1002]
    assert analysis["Run_ID"].tolist() == results["Run_ID"].tolist()
    assert len(results) == len(analysis) == 3
    assert not any(column.startswith("Event_Log_Raw_") for column in analysis)


def test_rm_and_add_fallback_are_analysis_only() -> None:
    runs = pd.DataFrame({"Run ID": [1]})
    logs = pd.DataFrame(
        {
            "run_id": [1, 1, 1],
            "log_no": [1, 2, 3],
            "time": ["00:00:00", "00:10:00", "00:20:00"],
            "event": [
                "Scrap basket added with: No1 Bundles: 48t; No1 Bundles: 41t",
                "Additions: Lime: 900kg; Lime: 530kg",
                "Tapping complete",
            ],
        }
    )
    results, analysis = build_excel_frames(runs, logs)
    assert not any(column.startswith("RM_") for column in results)
    assert not any(column.startswith("ADD_") for column in results)
    assert analysis.loc[0, "RM_Source"] == "EVENT_LOG_RECONSTRUCTED"
    assert analysis.loc[0, "RM_Reconstructed_No1_Bundles_t"] == 89
    assert analysis.loc[0, "ADD_Source"] == "EVENT_LOG_RECONSTRUCTED"
    assert analysis.loc[0, "ADD_Reconstructed_Lime_kg"] == 1430


def test_analysis_run_1021_cost_and_electrode_regression() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1021],
            "Process Type": ["Electric Arc Furnace"],
            "Time": [70],
            "Cost Breakdown > Power": [354.5],
            "Cost Breakdown > Scrap": [70],
            "Cost Breakdown > Additions": [16150],
            "Cost Breakdown > Other consumables": [18968.24],
            "Total Cost": [35771],
        }
    )
    logs = pd.DataFrame(
        {
            "run_id": [1021, 1021, 1021],
            "log_no": [1, 2, 3],
            "time": ["00:34:54", "00:46:22", "01:10:20"],
            "event": [
                "Electrode breakage",
                "Broken electrode has been replaced",
                "Tapping complete",
            ],
        }
    )
    analysis = build_analysis_dataframe(runs, logs).iloc[0]
    assert analysis["Corrected_Additions_Cost_USD"] == pytest.approx(354.5)
    assert analysis["Corrected_Scrap_Cost_USD"] == pytest.approx(16150)
    assert analysis["Corrected_Power_Cost_USD"] == pytest.approx(18968.24)
    assert analysis["Other_Consumables_USD"] == pytest.approx(298.26)
    assert analysis["Electrode_Break_Count"] == 1
    assert analysis["Electrode_Total_Downtime_sec"] == 688


def test_small_cost_residual_is_explicit_rounding_adjustment() -> None:
    runs = pd.DataFrame(
        {
            "Run ID": [1],
            "Process Type": ["Electric Arc Furnace"],
            "Time": [42],
            "Cost Breakdown > Power": [100],
            "Cost Breakdown > Scrap": [42],
            "Cost Breakdown > Additions": [200],
            "Cost Breakdown > Other consumables": [300.31],
            "Total Cost": [600],
        }
    )
    row = build_analysis_dataframe(runs, pd.DataFrame()).iloc[0]
    assert row["Other_Consumables_Raw_Residual"] == pytest.approx(-0.31)
    assert row["Other_Consumables_USD"] == 0
    assert row["Cost_Rounding_Adjustment_USD"] == pytest.approx(-0.31)


def test_public_raw_event_chunks_preserve_all_text() -> None:
    event = "Y" * 65_000
    results = build_results_dataframe(
        pd.DataFrame({"Run ID": [1]}),
        pd.DataFrame(
            {"run_id": [1], "log_no": [1], "time": ["00:00:00"], "event": [event]}
        ),
    )
    columns = sorted(
        (column for column in results if column.startswith("Event_Log_Raw_")),
        key=lambda value: int(value.rsplit("_", 1)[1]),
    )
    assert len(columns) == 3
    assert "".join(str(results.loc[0, column]) for column in columns) == f"00:00:00,{event}"
    assert all(len(str(results.loc[0, column])) <= 30_000 for column in columns)
